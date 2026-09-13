#!/usr/bin/env python3
"""qa.py 回归自测：改 qa.py 之后必跑。

在临时目录里造三份 deck，对闸门一的行为做逐项断言：

  broken.pptx   带病稿：20+ 病种，覆盖全部检查类别 → 期望 exit 1 且逐类命中
  small.pptx    10×5.625 默认画布 → 期望 [画布] 告警（pitfalls #9 自动化）
  clean.pptx    node + pptxgenjs 生成的干净稿 → 期望 exit 0 且零告警
                （缺 node / pptxgenjs 环境则跳过并注明）

历史教训（本文件存在的理由）：劈词检查曾因 jieba 未 initialize 静默失效了
很多次运行——检查会死，而且死了没有任何症状。所以每次动 qa.py 都要跑这里。

用法：python scripts/selftest.py
退出码：0 全过，1 有断言失败（失败时打印 broken 稿的完整 qa 输出供排查）。
"""
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
QA = HERE / "qa.py"

try:  # 裸机上没有 jieba 时，劈词断言要明确跳过而不是失败
    import jieba  # noqa: F401
    _HAS_JIEBA = True
except Exception:
    _HAS_JIEBA = False

CLEAN_JS = r"""
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
p.layout = "LAYOUT_WIDE";
p.title = "selftest clean";
const BG = "FFFFFF", PRIMARY = "1F3A5F", TEXT = "1A1A1A",
      MUTED = "6B7280", ACCENT = "9B2226", SANS = "微软雅黑";
const edge = () => ({ color: "111111", width: 3 });
const hard = (color = "111111", offset = 7) =>
  ({ type: "outer", color, blur: 0.0001, offset, angle: 45, opacity: 1.0 });
const lines = (slide, arr, opts) =>
  slide.addText(arr.map((t, i) => ({ text: t,
    options: i < arr.length - 1 ? { breakLine: true } : {} })), opts);
const s1 = p.addSlide();
s1.background = { color: BG };
lines(s1, ["迁移成本三年降 58%", "结论先行的页标题"], { x: 0.5, y: 0.5, w: 6, h: 1.6,
  margin: 0, fontFace: SANS, fontSize: 40, bold: true, color: TEXT, paraSpaceAfter: 4 });
s1.addShape(p.shapes.ROUNDED_RECTANGLE,
  { x: 0.5, y: 2.2, w: 6, h: 2.6, fill: { color: PRIMARY }, rectRadius: 0.14 });
lines(s1, ["主格标题", "单行安全容量估算：框宽乘七十二除字号。"], {
  x: 0.8, y: 2.5, w: 5.6, h: 0.9, margin: 0, fontFace: SANS, fontSize: 14,
  color: "FFFFFF", paraSpaceAfter: 2 });
s1.addShape(p.shapes.RECTANGLE,
  { x: 12.83 - 4.0, y: 2.2, w: 4.0, h: 0.6, fill: { color: "FFFFFF" },
    line: edge(), shadow: hard() });
s1.addText("KPI", { x: 9.23, y: 2.25, w: 3.6, h: 0.5, margin: 0,
  fontFace: SANS, fontSize: 16, bold: true, color: TEXT });
s1.addChart(p.charts.BAR, [{ name: "迁移成本",
  labels: ["2024", "2025", "2026"], values: [100, 62, 31] }], {
  x: 0.5, y: 5.0, w: 6, h: 2.2, barDir: "col", barGrouping: "clustered",
  varyColors: true, chartColors: [PRIMARY, PRIMARY, ACCENT],
  showValue: true, dataLabelPosition: "outEnd",
  dataLabelFontSize: 12, dataLabelFontFace: SANS, dataLabelColor: TEXT,
  catAxisLabelFontSize: 12, catAxisLabelFontFace: SANS, catAxisLabelColor: MUTED });
p.writeFile({ fileName: "clean_deck.pptx" }).then(() => console.log("ok"));
"""


def _run_qa(pptx: Path, *flags):
    r = subprocess.run([sys.executable, str(QA), str(pptx), *flags],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    return r.returncode, (r.stdout or "") + (r.stderr or "")


def build_broken(path: str):
    """带病稿。每个病种在注释里标了断言名，改病种时同步改 main() 的断言表。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from PIL import Image

    # 图片供压图 / band 几何用：图片与图表曾让 fill 访问裸崩，必须有它们常驻
    Image.new("RGB", (4, 3), (200, 0, 0)).save(str(Path(path).parent / "dot.png"))

    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    INK = RGBColor(0x1A, 0x1A, 0x1A)
    GRAY = RGBColor(0x8A, 0x8A, 0x8A)
    DARK = RGBColor(0x33, 0x33, 0x33)
    NOTE = RGBColor(0x6B, 0x72, 0x80)
    BAND = RGBColor(0xDC, 0xE7, 0xEF)
    BLUE = RGBColor(0x16, 0x3A, 0x63)

    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]

    def rect(slide, x, y, w, h, color, rot=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False
        if rot is not None:
            sh.rotation = rot
        return sh

    def text(slide, x, y, w, h, paras, zero_margin=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if zero_margin:
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (t, pt, color, font, bold) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = t
            r.font.size = Pt(pt)
            r.font.name = font
            r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color
        return tb

    # ---- S1：对比度 / 字号 / 混合字号 / 禁则 / 超宽 / 重叠 / 软换行 / 旋转 ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.33, 7.5, WHITE)
    text(s, 1, 1, 4, 0.6, [("看不见的我", 18, WHITE, "微软雅黑", False)])
    text(s, 1, 1.8, 4, 0.4, [("十一磅的小字", 11, DARK, "微软雅黑", False)])

    def _raw_shadow(shape, blur_rad):
        spPr = shape._element.spPr
        eff = spPr.makeelement(qn("a:effectLst"), {})
        shdw = eff.makeelement(qn("a:outerShdw"),
                               {"blurRad": str(blur_rad), "dist": "38100", "dir": "2700000"})
        clr = shdw.makeelement(qn("a:srgbClr"), {"val": "000000"})
        shdw.append(clr)
        eff.append(shdw)
        spPr.append(eff)

    sh_hard = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.3), Inches(5.2),
                                 Inches(1.5), Inches(0.6))
    _raw_shadow(sh_hard, 1)          # 硬阴影（blurRad 1）
    sh_soft = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.3), Inches(6.1),
                                 Inches(1.5), Inches(0.6))
    _raw_shadow(sh_soft, 101600)     # 柔和阴影（8pt）——同页混 shadows 不应互相掩护
    text(s, 5, 1, 6, 1.5, [("一眼看全", 40, INK, "微软雅黑", True),
                           ("逐字读有三个代价，一眼看全需要三块基石。", 14, INK, "微软雅黑", False)],
         zero_margin=True)  # 断言：不产生「溢出」误报
    text(s, 5, 3, 4, 1.2, [("混合字号框", 40, INK, "微软雅黑", True),
                           ("这一段正文故意写得特别长，就是为了把逐段计算的行高撑到超过框体，" * 2,
                            14, INK, "微软雅黑", False)],
         zero_margin=True)  # 断言：真溢出命中
    text(s, 1, 2.6, 4, 0.8, [("他说明了三件事", 14, INK, "微软雅黑", False),
                             ("，以及一个代价", 14, INK, "微软雅黑", False)])
    text(s, 1, 3.6, 2, 1.5, [("这是一个非常非常长的句子远远超过了两英寸的窄栏宽度限制",
                              14, INK, "微软雅黑", False)])
    text(s, 8, 1, 2.5, 0.8, [("甲框", 14, INK, "微软雅黑", False)])
    text(s, 8.3, 1.2, 2.5, 0.8, [("乙框", 14, INK, "微软雅黑", False)])
    text(s, 8, 3, 4, 0.5, [("次级灰字压白底", 14, GRAY, "微软雅黑", False)])
    # <a:br/> 软换行：两段各约六成宽，合计超宽——若 \v 被当字符算就会误报行超宽
    tb = text(s, 5, 4.6, 6, 1.0, [("br软换行第一段可以到八成宽不折行没问题啊", 14, INK, "微软雅黑", False)],
              zero_margin=True)
    # 公式残留：未转换的 LaTeX 显示为乱码文本，必须判硬伤（inject_math.py 漏网防线）
    text(s, 8.5, 4.6, 3, 0.6, [("$$\\frac{a}{b}$$", 14, INK, "微软雅黑", False)])
    # emoji 当图标：反 AI 清单明令禁止
    text(s, 8.5, 5.4, 3, 0.5, [("🚀 快速上线", 14, INK, "微软雅黑", False)])
    # 旋转窄框长文本：曾按未旋转宽度判行超宽 147% 硬伤（e2e 实测误报），现在必须跳过
    tb_rot = text(s, 2, 5.6, 0.45, 1.8,
                  [("旋转竖排文字长度远超窄框宽度用于验证不误报", 14, INK, "微软雅黑", False)])
    tb_rot.rotation = 270
    p0 = tb.text_frame.paragraphs[0]
    p0._p.append(p0._p.makeelement(qn("a:br"), {}))
    r2 = p0.add_run()
    r2.text = "第二段同样不该报超宽"
    r2.font.size, r2.font.name, r2.font.color.rgb = Pt(14), "微软雅黑", INK
    rect(s, 2, 5.2, 8, 0.25, BAND, rot=6)  # 断言：旋转提醒

    # ---- S2：同页眉 + 信息密度 + 页脚带 ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.33, 7.5, WHITE)
    text(s, 0.5, 0.3, 3, 0.5, [("第二章 背景", 14, INK, "微软雅黑", False)])
    text(s, 0.5, 1.2, 6, 5, [("本页信息密度超标的示例段落。" * 22, 14, INK, "微软雅黑", False)],
         zero_margin=True)
    rect(s, 0, 6.9, 13.33, 0.5, BAND)

    # ---- S3：同页眉 + 行尾禁则 + 非安全字体 + 页脚带 ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.33, 7.5, WHITE)
    text(s, 0.5, 0.3, 3, 0.5, [("第二章 背景", 14, INK, "微软雅黑", False)])
    text(s, 0.5, 1.2, 4, 0.5, [("他引用了一句话「", 14, INK, "微软雅黑", False)])
    text(s, 0.5, 2.0, 4, 0.5, [("非安全字体示例", 14, INK, "思源黑体", False)])
    text(s, 5.5, 2.0, 4.5, 0.5, [("中文排进拉丁字体", 14, INK, "Arial Black", False)])
    sh_hash = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(5.8),
                                 Inches(2), Inches(0.5))
    spPr = sh_hash._element.spPr
    ln = spPr.makeelement(qn("a:ln"), {"w": "19050"})
    solid = ln.makeelement(qn("a:solidFill"), {})
    clr = solid.makeelement(qn("a:srgbClr"), {"val": "#FF0000"})  # 带 # 的颜色会坏文件
    solid.append(clr)
    ln.append(solid)
    spPr.append(ln)
    tb = text(s, 5, 2.0, 5, 0.8, [("缩到一半的十六磅字", 16, INK, "微软雅黑", False)])
    tfs = tb.text_frame
    tfs._txBody.bodyPr.append(
        tfs._txBody.bodyPr.makeelement(qn("a:normAutofit"), {"fontScale": "50000"}))
    rect(s, 0, 6.9, 13.33, 0.5, BAND)

    # ---- S4：同页眉 + 劈词 + 孤行 + 页眉带 ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.33, 7.5, WHITE)
    text(s, 0.5, 0.3, 3, 0.5, [("第二章 背景", 14, INK, "微软雅黑", False)])
    text(s, 0.5, 1.2, 4, 1.0, [("这是渐进式算", 14, INK, "微软雅黑", False),
                               ("法介绍的第二段", 14, INK, "微软雅黑", False)])
    text(s, 5, 1.2, 4, 1.2, [("第一段有十几个字撑满这一行的宽度啊", 14, INK, "微软雅黑", False),
                             ("尾", 14, INK, "微软雅黑", False)])
    rect(s, 0, 0, 13.33, 0.4, BAND)

    # ---- S5：同页眉 + 12pt 脚注 + 未设色压深底 + autofit + 表格小字 + 填充文本框色带 ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.33, 7.5, WHITE)
    text(s, 0.5, 0.3, 3, 0.5, [("第二章 背景", 14, INK, "微软雅黑", False)])
    text(s, 0.5, 6.8, 6, 0.4, [("来源：示例数据，仅用于自测", 12, NOTE, "微软雅黑", False)])
    rect(s, 7, 2.5, 5.5, 1.2, BLUE)
    text(s, 7.2, 2.7, 5.1, 0.8, [("这行字忘了写颜色，继承主题近黑", 14, None, "微软雅黑", False)])
    tb = text(s, 7, 4.0, 5.5, 0.8, [("这个框配置了自动缩字，渲染字号会变", 14, INK, "微软雅黑", False)])
    tfx = tb.text_frame
    tfx._txBody.bodyPr.append(tfx._txBody.bodyPr.makeelement(qn("a:normAutofit"), {}))
    gfx = s.shapes.add_table(2, 3, Inches(0.5), Inches(2.5), Inches(6), Inches(1.6))
    tbl = gfx.table
    tbl.cell(0, 0).merge(tbl.cell(0, 1))  # 断言：合并格路径不崩、小字命中
    tfc = tbl.cell(0, 0).text_frame
    tfc.text = "表内小字"
    rc = tfc.paragraphs[0].runs[0]
    rc.font.size, rc.font.name, rc.font.color.rgb = Pt(10), "微软雅黑", INK
    tfc2 = tbl.cell(1, 1).text_frame
    tfc2.text = "表内正常字"
    rc2 = tfc2.paragraphs[0].runs[0]
    rc2.font.size, rc2.font.name, rc2.font.color.rgb = Pt(14), "微软雅黑", INK
    band_tb = s.shapes.add_textbox(Inches(0), Inches(6.9), Inches(13.33), Inches(0.5))
    band_tb.fill.solid()
    band_tb.fill.fore_color.rgb = BAND  # 断言：TEXT_BOX 填充色带也要被识别
    # 回归：band 几何的图片（曾让 fill 访问崩溃）、压图文字、图表（_bg_under 曾崩）
    s.shapes.add_picture(str(Path(path).parent / "dot.png"),
                         Inches(0), Inches(6.9), Inches(13.33), Inches(0.5))
    s.shapes.add_picture(str(Path(path).parent / "dot.png"),
                         Inches(7), Inches(5.4), Inches(5.5), Inches(1.3))
    text(s, 7.4, 5.7, 4.7, 0.6, [("压在图片上的字", 14, INK, "微软雅黑", False)])
    cd = CategoryChartData()
    cd.categories = ["a", "b"]
    cd.add_series("s", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                       Inches(0.5), Inches(4.6), Inches(5.5), Inches(2.0), cd)
    pie = CategoryChartData()
    pie.categories = ["一", "二", "三", "四", "五", "六"]
    pie.add_series("p", (30, 22, 16, 12, 10, 10))
    s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(5.6), Inches(4.3),
                       Inches(2.6), Inches(2.2), pie)  # 断言：饼图 6 块告警

    prs.save(path)


def build_small(path: str):
    """10×5.625 默认画布稿：断言 [画布] 告警。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "默认画布上的标题"
    r.font.size, r.font.name = Pt(28), "微软雅黑"
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    prs.save(path)


def _node_env() -> dict:
    """temp 目录没有 node_modules 祖先链，把 cwd 与技能目录沿途的 node_modules
    全部注入 NODE_PATH（npm 全局目录也一并附上）。"""
    roots: list[str] = []
    for base in (Path.cwd(), HERE):
        for p in (base, *base.parents):
            nm = p / "node_modules"
            if nm.is_dir() and str(nm) not in roots:
                roots.append(str(nm))
    try:
        npm_root = subprocess.run(["npm", "root", "-g"], capture_output=True,
                                  text=True).stdout.strip()
    except Exception:
        npm_root = ""
    if npm_root:
        roots.append(npm_root)
    env = dict(os.environ)
    env["NODE_PATH"] = os.pathsep.join(roots + [env.get("NODE_PATH", "")])
    return env


def build_clean(td: str) -> bool:
    """node + pptxgenjs 干净稿；环境不可用返回 False（跳过而非失败）。"""
    js = Path(td) / "build_clean.js"
    js.write_text(CLEAN_JS, encoding="utf8")
    env = _node_env()
    try:
        r = subprocess.run(["node", str(js)], capture_output=True, text=True,
                           cwd=td, env=env)
    except FileNotFoundError:
        return False
    return r.returncode == 0 and (Path(td) / "clean_deck.pptx").exists()


def main() -> int:
    failures: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        bp = Path(td) / "broken.pptx"
        build_broken(str(bp))
        code, out = _run_qa(bp)
        hard_lines = [l for l in out.splitlines() if "溢出" in l]

        checks = [
            ("带病稿 exit 1", code == 1),
            ("对比度硬伤（白字白底）", "对比度 1.0:1" in out),
            ("对比度告警（灰字白底）", "低于 4.5" in out),
            ("字号下限逐 run（11pt）", "字号 11pt" in out),
            ("表格小字命中（10pt）", "字号 10pt" in out),
            ("混合字号真溢出", "混合字号框" in "".join(hard_lines)),
            ("混合字号装得下不误报", "一眼看全" not in "".join(hard_lines)),
            ("行超宽", "行超宽" in out),
            ("软换行不误报行超宽", all("br软换行" not in l for l in hard_lines)),
            ("旋转文本框不判行超宽/溢出", all("旋转竖排" not in l for l in hard_lines)),
            ("行首禁则", "行首禁则" in out),
            ("行尾禁则", "行尾禁则" in out),
            ("劈词（jieba 已初始化）", "疑似劈词" in out),
            ("孤行尾行", "孤行尾行" in out),
            ("文本框重叠", "文本框重叠" in out),
            ("无显式 RGB 压深底兜底", "无显式 RGB" in out),
            ("同页混硬/柔阴影互相不掩护", "混有硬阴影与柔和阴影" in out),
            ("颜色带 # 判硬伤", "颜色值带 #" in out),
            ("图片拉伸变形", "图片拉伸变形" in out),
            ("饼图超 5 块告警", "块 > 5" in out),
            ("中文排拉丁字体告警", "拉丁专用字体" in out),
            ("表格小字整表汇总（不刷屏）", "表格" in out and "个单元格字号" in out),
            ("自动缩字提醒（未记录比）", "自动缩字" in out),
            ("缩字后按实际字号判硬伤", "×缩字" in out),
            ("公式残留判硬伤", "公式未转换" in out),
            ("emoji 当图标告警", "emoji" in out),
            ("旋转提醒", "旋转" in out or "净空" in out),
            ("色带（含填充文本框）", "页有全宽页眉/页脚色带" in out),
            ("同页眉骨架", "同一页眉骨架" in out),
            ("载体断档", "视觉载体断档" in out),
            ("信息密度字当量", "字当量" in out),
            ("非安全字体", "跨平台安全名单" in out),
        ]
        if not _HAS_JIEBA:
            checks = [c for c in checks if c[0] != "劈词（jieba 已初始化）"]
            print("  ? jieba 不可用，劈词断言跳过（qa 报告会标『检查降级』）")
        code14, out14 = _run_qa(bp, "--min-pt", "14")
        checks.append(("--min-pt 14 生效（12pt 脚注命中）", "字号 12pt < 14pt" in out14))
        codeP, outP = _run_qa(bp, "--palette")  # 图片/图表曾让 palette 循环裸崩
        checks.append(("--palette 报告输出（含图片图表不崩）", "[palette]" in outP
                       and "Traceback" not in outP))

        # 畸形输入：截断文件必须优雅退出（有信息、无 Traceback）
        trunc = Path(td) / "trunc.pptx"
        trunc.write_bytes(bp.read_bytes()[: int(bp.stat().st_size * 0.6)])
        _, outT = _run_qa(trunc)
        checks.append(("截断文件优雅退出（无 Traceback）", "Traceback" not in outT))

        sp = Path(td) / "small.pptx"
        build_small(str(sp))
        _, outS = _run_qa(sp)
        checks.append(("10 英寸画布告警", "[画布]" in outS))

        if build_clean(td):
            codeC, outC = _run_qa(Path(td) / "clean_deck.pptx")
            # 降级（jieba/Pillow/字体目录缺失）会让退出码变 2，但不是 deck 的错、不算硬伤。
            checks.append(("干净稿无硬伤", "硬伤 0" in outC))
            checks.append(("干净稿零（非降级）告警", "告警 0" in outC))
            checks.append(("干净稿 exit 0/2（降级不误判）", codeC in (0, 2)))
        else:
            print("  ? node / pptxgenjs 不可用，跳过干净稿测试")

        # 渲染链回归（render.py 曾出过画幅比拉伸、Open 泄漏进程等 bug，防线必须覆盖）。
        # 无 PowerPoint 且无 LibreOffice 的机器上优雅跳过，不计失败。
        def build_twopage(path):
            from pptx import Presentation
            from pptx.util import Inches, Pt, Emu
            from pptx.dml.color import RGBColor
            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
            for i in range(2):
                s = prs.slides.add_slide(prs.slide_layouts[6])
                tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(10), Inches(1))
                r = tb.text_frame.paragraphs[0].add_run()
                r.text = f"渲染自测 第{i + 1}页"
                r.font.size, r.font.name = Pt(44), "微软雅黑"
                r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            prs.save(path)

        RENDER = HERE / "render.py"
        tp = Path(td) / "twopage.pptx"
        build_twopage(str(tp))
        rd = subprocess.run(
            [sys.executable, str(RENDER), str(tp), "st", "--contact"],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            cwd=str(td), timeout=300)
        rd_out = (rd.stdout or "") + (rd.stderr or "")
        if "找不到渲染器" in rd_out:
            print("  ? 无可用渲染器，跳过渲染断言")
        else:
            pngs = list((Path(td) / "preview").glob("st*.png"))
            checks.append(("渲染 2 页 + 拼图", rd.returncode == 0 and len(pngs) == 2
                           and (Path(td) / "preview" / "_contact_st.png").exists()))
            r2 = subprocess.run(
                [sys.executable, str(RENDER), str(tp), "stsel", "--pages", "2"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
                cwd=str(td), timeout=300)
            pngs2 = list((Path(td) / "preview").glob("stsel*.png"))
            checks.append(("--pages 2 只渲染 1 页", r2.returncode == 0 and len(pngs2) == 1))
            r3 = subprocess.run(
                [sys.executable, str(RENDER), str(tp), "st", "--pages", "2"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
                cwd=str(td), timeout=300)
            checks.append(("--pages 只重渲指定页（保留其他页与拼图）", r3.returncode == 0
                           and (Path(td) / "preview" / "st1.png").exists()
                           and (Path(td) / "preview" / "st2.png").exists()
                           and (Path(td) / "preview" / "_contact_st.png").exists()))

        # 动画链回归（inject_anim.py 的 timing XML 曾有一个让 PowerPoint 整树静默
        # 丢弃的 bug，只有 COM 验证能抓到——这条链必须常驻防线）。
        INJECT = HERE / "inject_anim.py"
        ANIM_JS = r"""
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
p.layout = "LAYOUT_WIDE";
const s1 = p.addSlide();
s1.addText("自动淡入", { objectName: "标题", x: 1, y: 0.8, w: 8, h: 1, fontSize: 36 });
s1.addText("论点一", { objectName: "论点一", x: 1, y: 2.2, w: 5, h: 0.8, fontSize: 20 });
s1.addShape("rect", { objectName: "hero", x: 8, y: 2, w: 3, h: 3, fill: { color: "2E5BFF" } });
const s2 = p.addSlide();
s2.addShape("rect", { objectName: "hero", x: 9, y: 1, w: 3.5, h: 4.5, fill: { color: "2E5BFF" } });
s2.addText("平滑页", { objectName: "标题二", x: 1, y: 0.8, w: 6, h: 1, fontSize: 32 });
p.writeFile({ fileName: "__OUT__" }).then(() => console.log("ok"));
"""

        def node_build(js_code: str, path) -> bool:
            js = Path(td) / f"build_{abs(hash(js_code)) % 99999}.js"
            js.write_text(js_code.replace("__OUT__", str(path).replace("\\", "/")),
                          encoding="utf8")
            try:
                r = subprocess.run(["node", str(js)], capture_output=True, text=True,
                                   cwd=td, env=_node_env(), timeout=120)
            except FileNotFoundError:
                return False
            return r.returncode == 0 and Path(path).exists()

        def build_anim_deck(path):
            return node_build(ANIM_JS, path)

        ANIM_JSON = {
            "transitions": {"2": "平滑"},
            "pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "自动", "时长": 0.9},
                {"形状": "论点一", "效果": "出现", "触发": "之后"},
                {"形状": "hero", "效果": "浮入", "触发": "之后", "方向": "自底部"},
            ]},
        }
        if build_anim_deck(ap := str(Path(td) / "anim.pptx")):
            sj = Path(td) / "anim.json"
            sj.write_text(json.dumps(ANIM_JSON, ensure_ascii=False), encoding="utf8")
            ra = subprocess.run([sys.executable, str(INJECT), ap, str(sj)],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=300)
            a_out = (ra.stdout or "") + (ra.stderr or "")
            checks.append(("动画注入两页（timing+切换）", "已注入 2 页动画" in a_out
                           and "Traceback" not in a_out))
            if ra.returncode == 0 and "⚠" not in a_out:
                checks.append(("COM 三层验证全绿（逐条一致+往返一致）",
                               "逐条一致" in a_out and "往返校验一致" in a_out))
                # 只配切换、没有对象动画的页（本例 S2 的 morph）也必须进验证，
                # 否则切换读回错值会静默漏过（曾是 pages 只含对象动画页的旧行为）
                checks.append(("仅配切换的页也逐条验证切换（S2 平滑）",
                               "S2 切换 平滑 读回一致" in a_out))
            elif "⚠" in a_out:
                print("  ? 本机无 PowerPoint COM，动画验证断言降级为跳过")
            rg = subprocess.run([sys.executable, str(INJECT), ap, str(sj)],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=300)
            checks.append(("双重注入防护", rg.returncode != 0 and "已有动画" in (rg.stdout or "") + (rg.stderr or "")))
            codeA, outA = _run_qa(Path(ap))
            checks.append(("注入后 qa 干净（含 morph 同名配对）",
                           codeA == 0 and "不存在的形状 id" not in outA and "退化为普通淡入" not in outA))

            dang = Path(td) / "dangling.pptx"
            with zipfile.ZipFile(ap) as zin, \
                    zipfile.ZipFile(dang, "w", zipfile.ZIP_DEFLATED) as zout:
                for info in zin.infolist():
                    data = zin.read(info.filename)
                    if info.filename == "ppt/slides/slide1.xml":
                        data = re.sub(rb'spid="\d+"', b'spid="999"', data, count=1)
                    zout.writestr(info.filename, data)
            codeD, outD = _run_qa(dang)
            checks.append(("动画 spid 悬空判硬伤", codeD == 1 and "不存在的形状 id" in outD))
        else:
            print("  ? node / pptxgenjs 不可用，跳过动画链断言")

        # 原始XML 逃生舱（inject_anim.py --raw）：默认必须关闭，三重闸门必须挡住，
        # 正路径必须标 [降级] 且退出码非 0——手写 XML 没有逐条断言可做，不能报全绿。
        RAW_SRC = str(Path(td) / "rawsrc.pptx")
        if build_anim_deck(RAW_SRC):
            seed = Path(td) / "rawseed.json"
            seed.write_text(json.dumps({"pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "点击"}]}}, ensure_ascii=False), encoding="utf8")
            subprocess.run([sys.executable, str(INJECT), RAW_SRC, str(seed), "--no-verify"],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=300)
            txml = re.search(r"<p:timing>.*?</p:timing>",
                             zipfile.ZipFile(RAW_SRC).read("ppt/slides/slide1.xml")
                             .decode("utf-8"), re.S).group(0)
            GOOD = {"XML": txml, "原因": "回归：菜单外的效果", "用户已同意": True}

            def raw_case(name, spec, flags):
                p = Path(td) / f"raw_{name}.pptx"
                p.write_bytes(Path(RAW_SRC).read_bytes())      # 已带动画，靠 --replace 覆盖
                j = Path(td) / f"raw_{name}.json"
                j.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf8")
                before = p.read_bytes()
                r = subprocess.run([sys.executable, str(INJECT), str(p), str(j), *flags],
                                   capture_output=True, text=True, encoding="utf-8",
                                   errors="replace", timeout=300)
                return (r.returncode, (r.stdout or "") + (r.stderr or ""),
                        p.read_bytes() == before)

            def raw_spec(entry, extra=None):
                d = {"原始XML": {"1": entry}}
                if extra:
                    d.update(extra)
                return d

            co_ok, out_ok, _ = raw_case("ok", raw_spec(GOOD), ["--replace", "--raw"])
            checks.append(("原始XML 正路径：注入 + 标 [降级] + 提示逃生舱",
                           "逃生舱" in out_ok and "[降级]" in out_ok
                           and "没写「断言」" in out_ok and "Traceback" not in out_ok))
            checks.append(("原始XML 正路径退出码非 0（降级不报全绿）", co_ok == 2))

            # 「断言」清单：结构侧不需要 COM 也跑（这是逃生舱代价的解法）
            DEF = [{"形状": "标题", "presetID": 10, "起始ms": 0, "时长ms": 500}]   # 种子是「淡入」
            ca_ok, out_a, _ = raw_case(
                "assertok", raw_spec(dict(GOOD, **{"断言": DEF})), ["--replace", "--raw"])
            checks.append(("断言：全对 → 逐条一致且不再算降级",
                           "逐条断言一致" in out_a and "[降级]" not in out_a
                           and ca_ok in (0, 2)))
            ca_bad, out_ab, _ = raw_case(
                "assertbad", raw_spec(dict(GOOD, **{"断言": [
                    dict(DEF[0], **{"起始ms": 1200})]})), ["--replace", "--raw"])
            checks.append(("断言：起始ms 写错 → 报不一致（不需 COM 也能抓）",
                           ca_bad == 1 and "≠ 声明 1200ms" in out_ab))
            ca_cov, out_ac, _ = raw_case(
                "assertcov", raw_spec(dict(GOOD, **{"断言": DEF + DEF})), ["--replace", "--raw"])
            checks.append(("断言：覆盖不全（效果数对不上）→ 报错",
                           ca_cov == 1 and "效果数 1 ≠ 声明 2" in out_ac))
            ca_unk, out_au, _ = raw_case(
                "assertunk", raw_spec(dict(GOOD, **{"断言": [
                    dict(DEF[0], **{"起始": 0})]})), ["--replace", "--raw"])
            checks.append(("断言：未知字段被拒（挡笔误）",
                           ca_unk != 0 and "未知字段 起始" in out_au))
            # P0-1 回归：raw 页的切换照常验（原先 raw 分支的 continue 把它一起跳过了）
            _ct, out_t, _ = raw_case(
                "rawtrans", raw_spec(dict(GOOD, **{"断言": DEF}),
                                     extra={"transitions": {"1": "平滑"}}),
                ["--replace", "--raw"])
            if "没有 PowerPoint COM" in out_t:
                print("  ? 本机无 COM，跳过 raw 页切换断言")
            else:
                checks.append(("raw 页的切换照常验（P0-1 回归）",
                               "S1 切换 平滑 读回一致" in out_t))
            cm_a, out_ma, _ = raw_case("assertmenu", {
                "原始XML": {"2": GOOD},
                "pages": {"1": [{"形状": "标题", "效果": "淡入", "触发": "自动", "断言": []}]}},
                ["--replace", "--raw"])
            checks.append(("断言：写到菜单页上被拒（只给原始XML 用）",
                           cm_a != 0 and "未知字段 断言" in out_ma))

            co_no, out_no, same_no = raw_case("noraw", raw_spec(GOOD), ["--replace"])
            checks.append(("原始XML 无 --raw 被拒且文件未动",
                           co_no != 0 and "逃生舱" in out_no and same_no))

            co_nc, out_nc, same_nc = raw_case(
                "noconsent", raw_spec({"XML": txml, "原因": "x"}), ["--replace", "--raw"])
            checks.append(("原始XML 缺「用户已同意」被拒",
                           co_nc != 0 and "用户已同意" in out_nc and same_nc))

            co_nr, out_nr, _ = raw_case("noroot", raw_spec(
                {"XML": '<p:sld xmlns:p="http://schemas.openxmlformats.org/'
                        'presentationml/2006/main"/>', "原因": "x",
                 "用户已同意": True}), ["--replace", "--raw"])
            checks.append(("原始XML 根元素非 timing 被拒",
                           co_nr != 0 and "只接受 <p:timing>" in out_nr))

            co_dg, out_dg, same_dg = raw_case("dangling", raw_spec(
                {"XML": txml.replace('spid="2"', 'spid="999"'), "原因": "x",
                 "用户已同意": True}), ["--replace", "--raw"])
            checks.append(("原始XML 悬空 spid 被拒（PowerPoint 会静默丢弃）",
                           co_dg != 0 and "不存在的形状 id" in out_dg and same_dg))

            co_pf, out_pf, _ = raw_case("prefix", raw_spec(
                {"XML": txml.replace("</p:timing>", "<p159:x/></p:timing>"), "原因": "x",
                 "用户已同意": True}), ["--replace", "--raw"])
            checks.append(("原始XML 用 slide 未声明前缀被拒",
                           co_pf != 0 and "命名空间前缀" in out_pf))

            co_bt, out_bt, _ = raw_case("both", raw_spec(GOOD, extra={"pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "点击"}]}}), ["--replace", "--raw"])
            checks.append(("原始XML 与 pages 同页冲突被拒（一页只能一条 timing）",
                           co_bt != 0 and "二选一" in out_bt))
        else:
            print("  ? node / pptxgenjs 不可用，跳过原始XML 逃生舱断言")

        # 动效质量规则回归：缓动 / 整页自动连播 / 时长 / 自动换片 / morph 配对点名
        if build_anim_deck(ap3 := str(Path(td) / "anim_quality.pptx")):
            def anim_run(deck, spec, *flags):
                j = Path(deck).with_suffix(".json")
                j.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf8")
                r = subprocess.run([sys.executable, str(INJECT), deck, str(j), *flags],
                                   capture_output=True, text=True, encoding="utf-8",
                                   errors="replace", timeout=300)
                return r.returncode, (r.stdout or "") + (r.stderr or "")

            def slide1(deck):
                with zipfile.ZipFile(deck) as z:
                    return z.read("ppt/slides/slide1.xml").decode("utf-8")

            # ① 缓动：默认「缓入缓出」写进效果级 cTn；显式「缓出」只写 decel
            c1, o1 = anim_run(ap3, {"pages": {"1": [
                {"形状": "标题", "效果": "浮入", "触发": "自动", "方向": "自底部"},
                {"形状": "论点一", "效果": "淡入", "触发": "之后", "缓动": "缓出"},
            ]}}, "--no-verify")
            x1 = slide1(ap3)
            checks.append(("缓动：默认「缓入缓出」写入效果级 cTn",
                           c1 in (0, 2) and 'accel="50000" decel="50000"' in x1))
            checks.append(("缓动：显式「缓出」只写 decel（无 accel）",
                           c1 in (0, 2) and 'decel="50000" fill="hold"' in x1
                           and not re.search(r'accel="50000" fill="hold"', x1)))
            checks.append(("注入报告列出每页缓动", "缓动 缓入缓出×1／缓出×1" in o1))

            # ② 合理的节拍（2 组、每组 2 个）→ qa 不该有节拍告警
            _c, outN = _run_qa(Path(ap3))
            checks.append(("合理节拍（2 组×2 个）qa 无节拍告警",
                           "节拍" not in outN and "处点击" not in outN))

            # ③「出现」是瞬时效果，加缓动必须被拒
            c3, o3 = anim_run(ap3, {"pages": {"1": [
                {"形状": "标题", "效果": "出现", "触发": "自动", "缓动": "缓入缓出"}]}},
                "--replace", "--no-verify")
            checks.append(("「出现」加缓动被拒（时长 0，加缓动无意义）",
                           c3 != 0 and "瞬时效果" in o3))

            # ④ 缓动名非法 → 报错并列出可选项
            c4, o4 = anim_run(ap3, {"pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "自动", "缓动": "平滑"}]}},
                "--replace", "--no-verify")
            checks.append(("缓动名非法报错并列出可选项",
                           c4 != 0 and "缓动「平滑」不在菜单里" in o4))

            # ⑤ 分组就是触发语义：点击开新组、之后并进上一组 → 注入报告按组打印
            _c5, o5 = anim_run(ap3, {"pages": {"1": [
                {"形状": "标题", "效果": "浮入", "触发": "自动", "方向": "自底部"},
                {"形状": "论点一", "效果": "浮入", "触发": "之后", "方向": "自底部"},
                {"形状": "hero", "效果": "擦入", "触发": "点击", "方向": "自左侧"}]}},
                "--replace", "--no-verify")
            checks.append(("注入报告按组打印（组1 自动：标题、论点一）",
                           "2 组（自动 1／点击 1）" in o5
                           and "组1 自动：标题、论点一" in o5))
            # 节拍过碎：3 组全是单对象 → qa 告警
            anim_run(ap3, {"pages": {"1": [
                {"形状": n, "效果": "出现", "触发": "点击"}
                for n in ["标题", "论点一", "hero"]]}}, "--replace", "--no-verify")
            code5, out5 = _run_qa(Path(ap3))
            checks.append(("单对象成组（节拍过碎）qa 告警",
                           code5 == 2 and "节拍过碎" in out5))
            # 组数过多：6 组 → qa 告警
            anim_run(ap3, {"pages": {"1": [
                {"形状": n, "效果": "出现", "触发": "点击"} for n in
                ["标题", "论点一", "hero", "标题", "论点一", "hero"]]}},
                "--replace", "--no-verify")
            code5b, out5b = _run_qa(Path(ap3))
            checks.append(("节拍过多（6 组）qa 告警", code5b == 2 and "6 个揭示节拍" in out5b))

            # ⑥ 时长：单条过短 / 过长 / 单个节拍超 8s
            anim_run(ap3, {"pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "自动", "时长": 0.15},
                {"形状": "论点一", "效果": "淡入", "触发": "之后", "时长": 3.5}]}},
                "--replace", "--no-verify")
            _c6, out6 = _run_qa(Path(ap3))
            checks.append(("单条过短（0.15s）告警", "短于 0.2s" in out6))
            checks.append(("单条过长（3.5s）告警", "长于 3s" in out6))
            anim_run(ap3, {"pages": {"1": [
                {"形状": n, "效果": "淡入", "触发": ("自动" if i == 0 else "之后"), "时长": 2.9}
                for i, n in enumerate(["标题", "论点一", "hero"])]}},
                "--replace", "--no-verify")
            _c7, out7 = _run_qa(Path(ap3))
            checks.append(("单个节拍超 8s 告警（按绝对起点算，不是各条求和）",
                           "第 1 个节拍要播 8.7s" in out7))

            # 平滑的「选项」：p159:morph 的 option（探针：byObject/byWord/byChar → 3954/3955/3956）
            _co, o_opt = anim_run(ap3, {
                "transitions": {"1": {"效果": "平滑", "选项": "按词"}},
                "pages": {"1": [{"形状": "标题", "效果": "淡入", "触发": "自动"}]}}, "--replace")
            checks.append(("平滑「选项」按词 → 写 option=byWord", 'option="byWord"' in slide1(ap3)))
            if "没有 PowerPoint COM" in o_opt:
                print("  ? 本机无 COM，跳过 平滑选项 读回断言")
            else:
                checks.append(("平滑「选项」按词 → COM 读回 EntryEffect 3955",
                               "EntryEffect 3955" in o_opt))
            c_bo, o_bo = anim_run(ap3, {
                "transitions": {"1": {"效果": "平滑", "选项": "按段"}}}, "--replace", "--no-verify")
            checks.append(("平滑「选项」非法被拒（列出可选项）",
                           c_bo != 0 and "平滑切换的「选项」只能是" in o_bo))

            # ⑦ 自动换片时间（advTm / advClick=0）必须被拦
            anim_run(ap3, {"transitions": {"1": "淡入"}, "pages": {"1": [
                {"形状": "标题", "效果": "淡入", "触发": "自动"}]}}, "--replace", "--no-verify")
            adv = str(Path(td) / "anim_advtm.pptx")
            with zipfile.ZipFile(ap3) as zin, zipfile.ZipFile(adv, "w", zipfile.ZIP_DEFLATED) as zout:
                for info in zin.infolist():
                    d = zin.read(info.filename)
                    if info.filename == "ppt/slides/slide1.xml":
                        d = d.replace(b"<p:transition",
                                      b'<p:transition advTm="3000" advClick="0"', 1)
                    zout.writestr(info.filename, d)
            _c8, out8 = _run_qa(Path(adv))
            checks.append(("自动换片时间（advTm/advClick=0）告警",
                           "会自己翻过去" in out8 and "advTm=3000" in out8))

            # ⑧ 放映设置：使用计时 / 展台循环
            kio = str(Path(td) / "anim_kiosk.pptx")
            with zipfile.ZipFile(ap3) as zin, zipfile.ZipFile(kio, "w", zipfile.ZIP_DEFLATED) as zout:
                for info in zin.infolist():
                    d = zin.read(info.filename)
                    if info.filename == "ppt/presentation.xml":
                        s = d.decode("utf-8")
                        if "<p:showPr" in s:
                            s = re.sub(r"<p:showPr\b",
                                       '<p:showPr useTimings="1" showType="kiosk"', s, count=1)
                        else:
                            s = s.replace("</p:presentation>",
                                          '<p:showPr useTimings="1" showType="kiosk"/>'
                                          "</p:presentation>")
                        d = s.encode("utf-8")
                    zout.writestr(info.filename, d)
            _c9, out9 = _run_qa(Path(kio))
            checks.append(("放映设置「使用计时 / 展台循环」告警",
                           "使用计时" in out9 and "展台" in out9))
        else:
            print("  ? node / pptxgenjs 不可用，跳过动效质量规则断言")

        # morph 未配对：告警必须点名两页各自独有的形状（不然一轮改不对）
        MORPH_JS = r"""
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
p.layout = "LAYOUT_WIDE";
const s1 = p.addSlide();
s1.addText("封面", { objectName: "封面标题", x: 1, y: 1, w: 8, h: 1, fontSize: 36 });
s1.addShape("rect", { objectName: "装饰块", x: 1, y: 3, w: 3, h: 2, fill: { color: "2E5BFF" } });
const s2 = p.addSlide();
s2.addText("正文", { objectName: "正文标题", x: 1, y: 1, w: 8, h: 1, fontSize: 32 });
s2.addShape("rect", { objectName: "hero", x: 8, y: 2, w: 3, h: 3, fill: { color: "2E5BFF" } });
p.writeFile({ fileName: "__OUT__" }).then(() => console.log("ok"));
"""
        if node_build(MORPH_JS, mp2 := str(Path(td) / "anim_morph.pptx")):
            mj = Path(td) / "anim_morph.json"
            mj.write_text(json.dumps({"transitions": {"2": "平滑"}, "pages": {"2": [
                {"形状": "正文标题", "效果": "淡入", "触发": "自动"}]}}, ensure_ascii=False),
                encoding="utf8")
            subprocess.run([sys.executable, str(INJECT), mp2, str(mj), "--no-verify"],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=300)
            _cm, outm = _run_qa(Path(mp2))
            checks.append(("morph 未配对点名两页独有形状",
                           "退化成普通淡入" in outm and "封面标题" in outm
                           and "正文标题" in outm and "hero" in outm))
        else:
            print("  ? node / pptxgenjs 不可用，跳过 morph 点名断言")

        # 动画纪律预算：6 连击 → qa 告警
        if build_anim_deck(ap2 := str(Path(td) / "anim_over.pptx")):
            oj = Path(td) / "anim_over.json"
            oj.write_text(json.dumps({"pages": {"1": [
                {"形状": f"论点一", "效果": "出现", "触发": "点击"},
                {"形状": "标题", "效果": "出现", "触发": "点击"},
                {"形状": "hero", "效果": "出现", "触发": "点击"},
                {"形状": "标题", "效果": "出现", "触发": "点击"},
                {"形状": "hero", "效果": "出现", "触发": "点击"},
                {"形状": "论点一", "效果": "出现", "触发": "点击"},
            ]}}, ensure_ascii=False), encoding="utf8")
            subprocess.run([sys.executable, str(INJECT), ap2, str(oj), "--no-verify"],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=300)
            codeO, outO = _run_qa(Path(ap2))
            checks.append(("节拍过多/过碎告警（6 连击 = 6 组）",
                           codeO == 2 and "6 个揭示节拍" in outO))
        else:
            print("  ? node / pptxgenjs 不可用，跳过动画预算断言")

        # 公式链回归（inject_math.py）：块级公式→图片 + 金额防误判 + 与动画混跑的管线顺序
        INJECT_MATH = HERE / "inject_math.py"
        MATH_JS = r"""
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
p.layout = "LAYOUT_WIDE";
const s1 = p.addSlide();
s1.addText("会员价 $5、$8 两档", { objectName: "decoy", x: 0.5, y: 0.4, w: 12, h: 0.6,
  margin: 0, fontSize: 18, color: "404040", fontFace: "微软雅黑" });
s1.addText("$$E=mc^2$$", { x: 4, y: 1.6, w: 5, h: 1.4, margin: 0, fontSize: 36, color: "1A1A1A" });
s1.addText("$$\\int_0^1 x^2\\,dx = \\frac{1}{3}$$", { x: 2, y: 3.6, w: 9, h: 1.6,
  margin: 0, fontSize: 32, color: "1A1A1A", align: "center" });
p.writeFile({ fileName: "__OUT__" }).then(() => console.log("ok"));
"""
        if node_build(MATH_JS, mp := str(Path(td) / "math.pptx")):
            rm = subprocess.run([sys.executable, str(INJECT_MATH), mp],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=300)
            m_out = (rm.stdout or "") + (rm.stderr or "")
            checks.append(("公式注入 2 条（金额诱饵未被误转）", rm.returncode == 0
                           and "已插入 2 条公式图片" in m_out and "疑似转义丢失" not in m_out))
            # 管线顺序：python-pptx 重存后再做动画注入，必须仍然全绿
            aj = Path(td) / "math_anim.json"
            aj.write_text(json.dumps({"pages": {"1": [
                {"形状": "decoy", "效果": "淡入", "触发": "自动"}]}},
                ensure_ascii=False), encoding="utf8")
            ra2 = subprocess.run([sys.executable, str(INJECT), mp, str(aj)],
                                 capture_output=True, text=True, encoding="utf-8",
                                 errors="replace", timeout=300)
            a2out = (ra2.stdout or "") + (ra2.stderr or "")
            if ra2.returncode == 0 and "⚠" not in a2out:
                checks.append(("公式→动画混跑管线（动画仍逐条一致）",
                               "逐条一致" in a2out and "往返校验一致" in a2out))
            elif "⚠" in a2out:
                print("  ? 无 COM，混跑动画验证降级跳过")
            else:
                checks.append(("公式→动画混跑管线（动画仍逐条一致）", False))
            codeM, outM = _run_qa(Path(mp))
            checks.append(("注入后 qa 干净（公式为图片不算残留）",
                           codeM == 0 and "公式未转换" not in outM))
        else:
            print("  ? node / pptxgenjs 不可用，跳过公式链断言")

        # 行内公式报错路径：图片模式不支持，必须明确报错而非静默
        from pptx import Presentation as _P
        from pptx.util import Pt as _Pt, Inches as _In
        _ip = Path(td) / "inline.pptx"
        _prs = _P()
        _s = _prs.slides.add_slide(_prs.slide_layouts[6])
        _tb = _s.shapes.add_textbox(_In(1), _In(1), _In(8), _In(1))
        _r = _tb.text_frame.paragraphs[0].add_run()
        _r.text = "质能方程 $E=mc^2$ 很美"
        _r.font.size, _r.font.name = _Pt(18), "微软雅黑"
        _prs.save(str(_ip))
        ri = subprocess.run([sys.executable, str(INJECT_MATH), str(_ip)],
                            capture_output=True, text=True, encoding="utf-8",
                            errors="replace", timeout=300)
        iout = (ri.stdout or "") + (ri.stderr or "")
        checks.append(("行内公式明确报错（不静默）", ri.returncode != 0 and "行内公式" in iout))

        # 图标链（inject_icons.py）：锚点定位 + 调色板 + 防重注 + 每页预算
        INJECT_ICONS = HERE / "inject_icons.py"
        ICON_JS = r"""
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
p.layout = "LAYOUT_WIDE";
const s1 = p.addSlide();
s1.addText("三件事", { x: 0.7, y: 0.5, w: 11.9, h: 0.8, margin: 0,
  fontFace: "微软雅黑", fontSize: 28, bold: true, color: "111111" });
[["卡一", 0.9], ["卡二", 5.4], ["卡三", 9.9]].forEach(([n, x]) => {
  s1.addShape("rect", { objectName: n, x: x, y: 2.4, w: 3.0, h: 3.0, fill: { color: "F7F7F7" } });
});
p.writeFile({ fileName: "__OUT__" }).then(() => console.log("ok"));
"""
        if node_build(ICON_JS, ipath := str(Path(td) / "icons.pptx")):
            def icon_run(entries, *extra):
                sp = Path(td) / f"icons_{abs(hash(str(entries))) % 9999}.json"
                sp.write_text(json.dumps({"pages": {"1": entries}}, ensure_ascii=False),
                              encoding="utf8")
                return subprocess.run(
                    [sys.executable, str(INJECT_ICONS), ipath, str(sp),
                     "--colors", "INK=111111,ACCENT=E4002B", *extra],
                    capture_output=True, text=True, encoding="utf-8",
                    errors="replace", timeout=300)
            two = [{"图标": "安全", "锚点": "卡一", "位置": "内部上方居中", "颜色": "ACCENT", "尺寸": 0.6},
                   {"图标": "数据", "锚点": "卡二", "位置": "内部居中", "颜色": "INK", "尺寸": 0.6}]
            ri = icon_run(two)
            iout = (ri.stdout or "") + (ri.stderr or "")
            checks.append(("图标注入 2 个 + 有效 PPI 验证", ri.returncode == 0
                           and "已插入 2 个图标" in iout and "有效 PPI" in iout))
            rg = icon_run(two)
            gout = (rg.stdout or "") + (rg.stderr or "")
            checks.append(("图标重复注入被拦", rg.returncode != 0 and "已有" in gout))
            codeI, outI = _run_qa(Path(ipath))
            checks.append(("注入后 qa 干净（图标计入载体）", codeI == 0))
            four = [{"图标": "安全", "锚点": "卡一", "位置": "内部居中", "颜色": "INK", "尺寸": 0.35},
                    {"图标": "数据", "锚点": "卡一", "位置": "内部居中", "颜色": "INK", "尺寸": 0.35},
                    {"图标": "火箭", "锚点": "卡二", "位置": "内部居中", "颜色": "INK", "尺寸": 0.35},
                    {"图标": "目标", "锚点": "卡三", "位置": "内部居中", "颜色": "INK", "尺寸": 0.35}]
            icon_run(four, "--replace")
            codeI4, outI4 = _run_qa(Path(ipath))
            checks.append(("每页图标预算告警（4 个）", codeI4 == 2 and "个图标" in outI4))
        else:
            print("  ? node / pptxgenjs 不可用，跳过图标链断言")

        # 素材管线（prep_assets.py）：清点报告 + 裁切/双色调输出
        PREP = HERE / "prep_assets.py"
        try:
            from PIL import Image as _Img
            adir = Path(td) / "assets"
            adir.mkdir()
            _Img.new("RGB", (1200, 900), (60, 90, 160)).save(adir / "p.jpg", quality=90)
            rp = subprocess.run([sys.executable, str(PREP), str(adir)],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=120)
            pout = (rp.stdout or "") + (rp.stderr or "")
            checks.append(("素材清点报告", "素材清点" in pout and "1200×900" in pout))
            aspec = Path(td) / "aspec.json"
            aspec.write_text(json.dumps({"p.jpg": {"裁切": "16:9", "目标宽": 6.0,
                                                   "双色调": ["111111", "E4002B"]}},
                                        ensure_ascii=False), encoding="utf8")
            rp2 = subprocess.run([sys.executable, str(PREP), str(adir), str(aspec)],
                                 capture_output=True, text=True, encoding="utf-8",
                                 errors="replace", timeout=180)
            outjpg = adir / "_ready" / "p.jpg"
            okp = rp2.returncode == 0 and outjpg.exists()
            if okp:
                with _Img.open(outjpg) as im:
                    okp = im.size[0] == 1800 and abs(im.size[0] / im.size[1] - 16 / 9) < 0.03
            checks.append(("素材调理（裁切 16:9 + 双色调，1800px）", okp))
        except Exception:
            checks.append(("素材调理（裁切 16:9 + 双色调，1800px）", False))

        # 有效 PPI：小图放大放置必须判硬伤
        try:
            from PIL import Image as _Img2
            from pptx import Presentation as _PP
            from pptx.util import Inches as _In2
            tiny = Path(td) / "tiny.png"
            _Img2.new("RGB", (60, 45), (200, 30, 30)).save(tiny)
            tpp = Path(td) / "tinydeck.pptx"
            pr = _PP()
            sl = pr.slides.add_slide(pr.slide_layouts[6])
            sl.shapes.add_picture(str(tiny), _In2(0.5), _In2(0.5),
                                  width=_In2(10), height=_In2(7.5))
            pr.save(str(tpp))
            cP, oP = _run_qa(tpp)
            checks.append(("有效 PPI 过低判硬伤", cP == 1 and "有效 PPI" in oP))
        except Exception:
            checks.append(("有效 PPI 过低判硬伤", False))

        # 正文几何符号（△□∠）→ 告警：中文字体里它们只有汉字四成大小，LLM 文本味（pitfalls #29）
        try:
            from pptx import Presentation as _PPG
            from pptx.util import Inches as _InG
            gpp = Path(td) / "glyph.pptx"
            prg = _PPG()
            prg.slide_width, prg.slide_height = _InG(13.333), _InG(7.5)
            slg = prg.slides.add_slide(prg.slide_layouts[6])
            slg.shapes.add_textbox(_InG(1), _InG(1), _InG(8), _InG(1)) \
                .text_frame.text = "求 △PAB 的面积，如图 □ABCD，∠A = 90°"
            prg.save(str(gpp))
            cGl, oGl = _run_qa(gpp)
            checks.append(("正文几何符号（△□∠）告警", cGl == 2 and "几何符号" in oGl))
        except Exception:
            checks.append(("正文几何符号（△□∠）告警", False))

        # 表格外框 vs 行高合计：外框比内容矮 → 真 PowerPoint 裁剪（WPS/LibreOffice 撑开，
        # 渲染预览看不见）；自洽的表格不得误报（pitfalls #30）
        try:
            from pptx import Presentation as _PPT
            from pptx.util import Inches as _InT
            tdp = Path(td) / "tblclip.pptx"
            prt = _PPT()
            prt.slide_width, prt.slide_height = _InT(13.333), _InT(7.5)
            slt = prt.slides.add_slide(prt.slide_layouts[6])
            shp = slt.shapes.add_table(5, 2, _InT(0.6), _InT(1.0), _InT(6.0), _InT(3.6))
            shp.height = _InT(1.0)          # 复现库坑：行高合计 3.6" 而外框只 1.0"
            prt.save(str(tdp))
            cTc, oTc = _run_qa(tdp)
            checks.append(("表格外框矮于行高合计判硬伤", cTc == 1 and "表格外框高" in oTc))

            tdp2 = Path(td) / "tblok.pptx"
            prt2 = _PPT()
            prt2.slide_width, prt2.slide_height = _InT(13.333), _InT(7.5)
            slt2 = prt2.slides.add_slide(prt2.slide_layouts[6])
            slt2.shapes.add_table(5, 2, _InT(0.6), _InT(1.0), _InT(6.0), _InT(3.6))
            prt2.save(str(tdp2))
            cTo, oTo = _run_qa(tdp2)
            checks.append(("表格外框自洽不误报", "表格外框高" not in oTo and "表格列宽合计" not in oTo))
        except Exception:
            checks.append(("表格外框矮于行高合计判硬伤", False))
            checks.append(("表格外框自洽不误报", False))

        # 图片型强调元素对比度：暗色图标压深底（图片不进文字对比度检查）必须判硬伤
        try:
            from PIL import Image as _Img3
            from pptx import Presentation as _PP3
            from pptx.util import Inches as _In3
            from pptx.dml.color import RGBColor as _RGB3
            icp = Path(td) / "darkicon.png"
            _Img3.new("RGB", (64, 64), (193, 59, 46)).save(icp)      # 印泥红 C13B2E
            dpp = Path(td) / "icontest.pptx"
            pr3 = _PP3()
            sl3 = pr3.slides.add_slide(pr3.slide_layouts[6])
            bgsh = sl3.shapes.add_shape(1, _In3(0), _In3(0), _In3(13.33), _In3(7.5))
            bgsh.fill.solid()
            bgsh.fill.fore_color.rgb = _RGB3(0x1D, 0x4E, 0x89)       # 蓝图蓝
            bgsh.line.fill.background()
            pic3 = sl3.shapes.add_picture(str(icp), _In3(2), _In3(2),
                                          width=_In3(1), height=_In3(1))
            pic3._element.nvPicPr.cNvPr.set("name", "icon:测试")
            pr3.save(str(dpp))
            cG, oG = _run_qa(dpp)
            checks.append(("图片型强调元素对比度兜底（红压蓝判硬伤）",
                           cG == 1 and "图形对比度" in oG))
            pic4 = sl3.shapes.add_picture(str(icp), _In3(4), _In3(2),
                                          width=_In3(1), height=_In3(1))
            pic4._element.nvPicPr.cNvPr.set("name", "figure:测试")
            pr3.save(str(dpp))
            cH, oH = _run_qa(dpp)
            checks.append(("figure:* 同样进图形对比度兜底", cH == 1 and "figure" in oH))
        except Exception:
            checks.append(("图片型强调元素对比度兜底（红压蓝判硬伤）", False))
            checks.append(("figure:* 同样进图形对比度兜底", False))

        # 插图链（inject_figure.py）：约束验算通过 + 违例中止不动文件 + 逃生舱标注
        INJECT_FIG = HERE / "inject_figure.py"
        try:
            from pptx import Presentation as _PPF
            from pptx.util import Inches as _InF
            figbase = Path(td) / "figbase.pptx"
            prf = _PPF()
            prf.slide_width, prf.slide_height = _InF(13.333), _InF(7.5)
            slf = prf.slides.add_slide(prf.slide_layouts[6])
            tb = slf.shapes.add_textbox(_InF(0.7), _InF(1.0), _InF(5.0), _InF(4.6))
            tb.name = "图区"
            tb.text_frame.text = ""
            prf.save(str(figbase))

            def fig_run(spec, *extra):
                deckn = Path(td) / f"fig_{abs(hash(str(spec))) % 9999}.pptx"
                deckn.write_bytes(figbase.read_bytes())
                spn = Path(td) / f"fig_{abs(hash(str(spec))) % 9999}.json"
                spn.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf8")
                return subprocess.run(
                    [sys.executable, str(INJECT_FIG), str(deckn), str(spn),
                     "--colors", "INK=111111,ACCENT=E4002B", *extra],
                    capture_output=True, text=True, encoding="utf-8",
                    errors="replace", timeout=420), deckn

            tri = {"点": [{"名": "A", "坐标": [0, 0]}, {"名": "B", "坐标": [4, 0]},
                          {"名": "C", "坐标": [0, 3]}],
                   "线段": [{"点": ["A", "B"]}, {"点": ["A", "C"]}, {"点": ["B", "C"]}],
                   "直角": [{"顶点": "A", "边": ["B", "C"]}]}
            good = {"pages": {"1": [{"锚点": "图区", "标签": "三角形", "图": dict(
                tri, **{"坐标": {"x": [-1, 5], "y": [-1, 4]}, "等比例": True,
                        "约束": [{"类型": "直角", "顶点": "A", "边": ["B", "C"]},
                                 {"类型": "长度", "线段": ["A", "B"], "长度": 4.0}]})}]}}
            rF, deckF = fig_run(good)
            fout = (rF.stdout or "") + (rF.stderr or "")
            checks.append(("插图注入 + 几何约束验算通过 + 有效 PPI", rF.returncode == 0
                           and "已插入 1 张插图" in fout and "约束 2 条验算通过" in fout
                           and "有效 PPI" in fout))
            codeF, outF = _run_qa(deckF)
            checks.append(("插图注入后 qa 干净（figure:* 计入载体）", codeF == 0))
            try:
                import io as _io
                from PIL import Image as _ImT
                transp = False
                for _s in _PPF(str(deckF)).slides:
                    for _sh in _s.shapes:
                        if (_sh.name or "").startswith("figure:"):
                            with _ImT.open(_io.BytesIO(_sh.image.blob)) as _im:
                                transp = _im.convert("RGBA").getchannel("A").getextrema()[0] == 0
                checks.append(("插图透明底（transparent 生效，不白底压深版）", transp))
            except Exception:
                checks.append(("插图透明底（transparent 生效，不白底压深版）", False))

            bad = {"pages": {"1": [{"锚点": "图区", "标签": "违例", "图": dict(
                tri, **{"约束": [{"类型": "等长", "对象": [["A", "B"], ["A", "C"]]}]})}]}}
            rB, deckB = fig_run(bad)
            bout = (rB.stdout or "") + (rB.stderr or "")
            nfig = sum(1 for s in _PPF(str(deckB)).slides for sh in s.shapes
                       if (sh.name or "").startswith("figure:"))
            checks.append(("约束超差中止且不写入文件", rB.returncode != 0
                           and "超差" in bout and nfig == 0))

            esc = {"pages": {"1": [{"锚点": "图区", "标签": "逃生舱", "图": {
                "坐标": {"x": [-1, 2], "y": [-1, 2]},
                "自定义代码": 'ax.plot([0, 1, 2], [0, 1, 0], color=c("ACCENT"), lw=2)'}}]}}
            rE, _deckE = fig_run(esc)
            eout = (rE.stdout or "") + (rE.stderr or "")
            checks.append(("逃生舱自定义代码可用且标注未受约束验算",
                           rE.returncode == 0 and "未受约束验算" in eout))
        except Exception:
            for nm in ("插图注入 + 几何约束验算通过 + 有效 PPI",
                       "插图注入后 qa 干净（figure:* 计入载体）",
                       "约束超差中止且不写入文件",
                       "逃生舱自定义代码可用且标注未受约束验算"):
                checks.append((nm, False))

        # 旋转净空（qa 的旋转形状 ↔ 文字分离轴检查）：几何直接单元测，不依赖渲染
        try:
            import importlib.util as _iluR
            _spR = _iluR.spec_from_file_location("qa_rot", str(HERE / "qa.py"))
            _qaR = _iluR.module_from_spec(_spR)
            _spR.loader.exec_module(_qaR)

            class _Box:
                def __init__(self, l, t, w, h, rot=0.0):
                    self.left, self.top = int(l * 914400), int(t * 914400)
                    self.width, self.height = int(w * 914400), int(h * 914400)
                    self.rotation = rot

            # 远距：旋转 4° 的宽条与下方 1 英寸外的文字 → 必须判"分离"，不能判相交
            far = _qaR._poly_relation(_qaR._corners(_Box(0, 0, 4, 1, -4)),
                                      _qaR._corners(_Box(0, 6, 4, 1)))
            # 近距：净空 0.05 英寸 → 报正的下界
            near = _qaR._poly_relation(_qaR._corners(_Box(0, 0, 4, 1)),
                                       _qaR._corners(_Box(0, 1.05, 4, 1)))
            # 相交：重叠 → 报负值（深度）
            hit = _qaR._poly_relation(_qaR._corners(_Box(0, 0, 4, 1)),
                                      _qaR._corners(_Box(2, 0.2, 4, 1)))
            checks.append(("旋转净空：远距判分离（分离轴优先，曾满屏误报）", far > 1))
            checks.append(("旋转净空：近距给出实测下界（≈0.05）", 0.0 < near < 0.1))
            checks.append(("旋转净空：相交报负值（重叠深度）", hit < 0))
        except Exception as e:
            print("  ? 旋转净空单元断言跳过：%s" % str(e)[:60])

        # 扫描件清洗（prep_assets.py 的「扫描件」）：白底 + 去斜 + 报告
        try:
            import importlib.util as _ilu
            from PIL import Image as _ImgS, ImageDraw as _DrawS
            sdir = Path(td) / "scans"
            sdir.mkdir()
            imS = _ImgS.new("RGB", (900, 650), (212, 206, 188))
            dS = _DrawS.Draw(imS)
            dS.rectangle([150, 300, 750, 330], fill=(28, 28, 32))
            dS.rectangle([150, 120, 380, 260], outline=(28, 28, 32), width=5)
            imS = imS.rotate(3.0, resample=_ImgS.BICUBIC, expand=True, fillcolor=(212, 206, 188))
            imS.save(sdir / "scan.jpg", quality=92)

            def _ink_ratio(p):
                # 用「墨迹 bbox 高/宽」而不是像素高：清洗后会被降采样/放大，像素高不可比
                g = _ImgS.open(p).convert("L")
                bx = g.point(lambda v: 255 if v < 110 else 0).getbbox()
                if not bx:
                    return 0.0
                return (bx[3] - bx[1]) / max(1, bx[2] - bx[0])

            r0 = _ink_ratio(sdir / "scan.jpg")
            sspec = Path(td) / "sspec.json"
            sspec.write_text(json.dumps({"scan.jpg": {"扫描件": True, "灰度": True,
                                                      "目标宽": 5.0}}, ensure_ascii=False),
                             encoding="utf8")
            rS = subprocess.run([sys.executable, str(PREP), str(sdir), str(sspec)],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=240)
            sout = (rS.stdout or "") + (rS.stderr or "")
            sout_img = sdir / "_ready" / "scan.jpg"
            ok_bg = False
            r1 = r0
            if sout_img.exists():
                o = _ImgS.open(sout_img).convert("RGB")
                px = o.getpixel((3, 3))
                ok_bg = min(px) > 245 and (max(px) - min(px)) <= 2
                r1 = _ink_ratio(sout_img)
            checks.append(("扫描件清洗：白底归一 + 报告逐项", "扫描件：" in sout and "白底归一" in sout
                           and ok_bg))
            if _ilu.find_spec("numpy") is not None:
                checks.append(("扫描件去斜有效（墨迹框高宽比变小）", "去斜" in sout and r1 < r0))
            else:
                checks.append(("扫描件去斜（缺 numpy，明确跳过不静默）", "去斜跳过" in sout))
        except Exception:
            checks.append(("扫描件清洗：白底归一 + 报告逐项", False))
            checks.append(("扫描件去斜有效（墨迹框高宽比变小）", False))

        # 裁切区域：一页多图取单张（比例写法 / 像素写法）
        try:
            import importlib.util as _ilu2
            _sp2 = _ilu2.spec_from_file_location("prep_assets_mod", str(PREP))
            pam = _ilu2.module_from_spec(_sp2)
            _sp2.loader.exec_module(pam)
            from PIL import Image as _ImR
            _src = _ImR.new("RGB", (400, 300), (255, 255, 255))
            r_frac = pam._crop_region(_src, [0.25, 0.5, 0.75, 1.0]).size
            r_px = pam._crop_region(_src, [100, 60, 300, 240]).size
            checks.append(("裁切区域取图（比例/像素两种写法）",
                           r_frac == (200, 150) and r_px == (200, 180)))
        except Exception:
            checks.append(("裁切区域取图（比例/像素两种写法）", False))

        # --- 报告分诊 / 严格旗标 / 降级不静默（本轮新增）---
        sys.path.insert(0, str(HERE))
        import qa as _qa
        rl = _qa._report_lines(
            ["[S3] 甲", "[S1] 乙"], ["[S2] 丙", "[交付] 丁"], ["jieba 未装"], 12, "真实字体")
        rtext = "\n".join(rl)
        sys.path.remove(str(HERE))
        checks.append(("报告：计数行在发现列表之前",
                       rtext.index("硬伤 2 · 告警 2") < rtext.index("✗")))
        checks.append(("报告：按页升序、全局垫后",
                       rtext.index("[S1]") < rtext.index("[S3]") < rtext.index("[交付]")))
        checks.append(("报告：多页插页分隔行",
                       "── 第 1 页 ──" in rtext and "── 第 3 页 ──" in rtext))
        checks.append(("报告：降级成块并计入计数",
                       "检查降级 1" in rtext and "[降级] jieba 未装" in rtext))

        cBad, oBad = _run_qa(bp, "--min-pt", "abc")
        checks.append(("qa：数值旗标给错即报错", cBad != 0 and "不是数字" in oBad))
        cUnk, oUnk = _run_qa(bp, "--bogus")
        checks.append(("qa：未知旗标报错非静默", cUnk != 0 and "未知参数" in oUnk))

        def _run_py(script, *a):
            r = subprocess.run([sys.executable, str(HERE / script), *a],
                               capture_output=True, text=True, encoding="utf-8", errors="replace")
            return r.returncode, (r.stdout or "") + (r.stderr or "")
        cF, oF = _run_py("inject_figure.py", "x.pptx", "y.json", "--dpi", "abc")
        checks.append(("inject_figure：数值旗标给错即报错", cF != 0 and "不是整数" in oF))
        cP, oP = _run_py("prep_assets.py", str(HERE), "不存在.json")
        checks.append(("prep_assets：规格读不了无 Traceback",
                       cP != 0 and "读不了" in oP and "Traceback" not in oP))

        for name, ok in checks:
            print(("  ✅ " if ok else "  ❌ ") + name)
            if not ok:
                failures.append(name)

        if code != 1 or failures:
            print("\n---------- broken 稿 qa 完整输出 ----------\n" + out)

    print(f"\n通过 {len(checks) - len(failures)}/{len(checks)}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
