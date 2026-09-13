#!/usr/bin/env python3
"""闸门一：pptx 程序化检查。

五组检查：

  A 几何/排版    溢出（逐段落真实行高，含表格单元格）、单行超宽、文本框重叠、
                出画布、字号下限（逐 run）、图片拉伸变形、画布尺寸、自动缩字、旋转提醒
  B 中文排版     禁则处理（行首/行尾禁则）、跨行劈词、孤行、软换行（<a:br/>）、
                中文排进拉丁专用字体（渲染回退成衬线体，pitfalls #21）
  C 对比度/密度  文字与下方不透明纯色底的 WCAG 对比度（压图/图表=底色未知，跳过）；
                无显式 RGB 压深底兜底；每页字数当量粗查（拉丁字母按半字计）
  D 交付合法性   OPC 结构、单段落多 pPr、堆叠柱状图标签、饼图块数、颜色带 #、
                硬/柔阴影混用、字体可移植性
  E deck 级      骨架复用、全宽色带（页眉+页脚，含填充文本框画的色带）、
                连续无图无图表的载体断档

文本宽度用 Pillow 读**真实字体度量**，不用估算。差别很大：微软雅黑的大写 A 是
0.70em、Arial Black 是 0.78em，而常用的单一估算常数 0.55em 会同时造成误报与漏报。
字体在本机找不到时回落估算，并单独告警（字体缺失会导致静默替换，见 pitfalls #1）。
Pillow 整个缺失时如实报「全部估算」，绝不谎称真实度量。

溢出按**逐段落**计算：每段用自己的字号、行距（line_spacing）、段前段后距，
并扣除文本框内边距——混合字号的框（40pt 标题 + 14pt 正文）按整框最大字号
算是系统性误报的来源。<a:br/> 软换行按真实断行计。

对比度只查「显式 RGB 文字 × 不透明纯色底」；压图、渐变、主题色继承查不了，
交给闸门二（pitfalls #20）。run 完全没写颜色时继承主题（通常近黑），
压深底等于看不见，单独兜底告警。

表格单元格与文本框走同一条 audit_text_frame 检查路径；单元格不进重叠检查
（相邻单元格共享边框，按重叠判会全误报）。

用法：python qa.py deck.pptx [--min-pt 12] [--max-chars 300] [--palette]
  --min-pt N      最小字号下限（默认 12；投屏场景建议 14）
  --max-chars N   每页字数当量告警阈值（默认 300；投屏建议自降到 150）
  --palette       输出每页填充色面积占比（近似，重叠不互斥），人工核对用色纪律

报告分「硬伤 / 告警 / 检查降级」三类：降级指某类检查因环境缺失整类没跑
（jieba 未装致劈词检查关闭、Pillow 缺席致真实度量回落、非 Windows 无字体目录），
**降级不等于通过**——报告里会单列成块，交付说明不得声称该类已通过。

退出码：0 干净，1 有硬伤，2 只有告警或存在检查降级。
"""
from __future__ import annotations

import re
import sys
import unicodedata
import zipfile
from functools import lru_cache
from itertools import combinations
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    sys.exit("缺依赖：python -m pip install python-pptx   （qa.py 的硬依赖）")

EMU = 914400
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F]")
# 中文字体里"只有汉字四成大小、笔画细一档"的几何/数学符号（实测 △ 0.39em/墨量165，
# 而 角 0.98em/1454；② ½ 是满尺寸的，不算）。正文里用它 = 一眼 LLM 文本味，pitfalls #29。
_SHAPE_GLYPH = re.compile("[\u25b3\u25b2\u25a1\u25a0\u25c7\u25c6\u25cb\u25cf"
                          "\u2220\u22a5\u2225\u2235\u2234\u224c\u2245\u223d]")
_MAX_XML = 64 * 1024 * 1024

# pptx 内的 XML 来自外部文件，按不可信输入处理：优先 defusedxml，
# 否则用标准库但拒绝 DTD / 实体声明并限制体积，避免实体扩展耗尽资源。
try:
    from defusedxml.ElementTree import fromstring as _xml_parse
except ImportError:  # pragma: no cover
    import xml.etree.ElementTree as _ET

    def _xml_parse(data: bytes):
        if len(data) > _MAX_XML:
            raise ValueError("XML 超过体积上限")
        head = data[:4096].lower()
        if b"<!doctype" in head or b"<!entity" in head:
            raise ValueError("拒绝含 DTD / 实体声明的 XML")
        return _ET.fromstring(data)

# ---------------------------------------------------------------- CLI

def _cli_val(flag: str, default, cast=float):
    """读「--flag 值」或「--flag=值」。给了值却解析不了就报错退出，绝不静默退回默认值。"""
    unit = "整数" if cast is int else "数字"
    args = sys.argv[1:]
    for i, a in enumerate(args):
        raw = None
        if a == flag:
            if i + 1 >= len(args):
                sys.exit(f"{flag} 缺少值（用法：{flag} <{unit}>）")
            raw = args[i + 1]
        elif a.startswith(flag + "="):
            raw = a.split("=", 1)[1]
        if raw is not None:
            try:
                return cast(raw)
            except ValueError:
                sys.exit(f"{flag} 的值 {raw!r} 不是{unit}")
    return default

MIN_PT = _cli_val("--min-pt", 12.0)
MAX_CHARS = _cli_val("--max-chars", 300, int)
WANT_PALETTE = "--palette" in sys.argv

# ---------------------------------------------------------------- 字体度量

_FONT_DIR = Path("C:/Windows/Fonts")
# family(小写) -> (常规文件, 粗体文件)
_FONT_FILES = {
    "微软雅黑": ("msyh.ttc", "msyhbd.ttc"),
    "microsoft yahei": ("msyh.ttc", "msyhbd.ttc"),
    "宋体": ("simsun.ttc", "simsun.ttc"),
    "simsun": ("simsun.ttc", "simsun.ttc"),
    "黑体": ("simhei.ttf", "simhei.ttf"),
    "simhei": ("simhei.ttf", "simhei.ttf"),
    "等线": ("Deng.ttf", "Dengb.ttf"),
    "dengxian": ("Deng.ttf", "Dengb.ttf"),
    "楷体": ("simkai.ttf", "simkai.ttf"),
    "仿宋": ("simfang.ttf", "simfang.ttf"),
    "arial": ("arial.ttf", "arialbd.ttf"),
    "arial black": ("ariblk.ttf", "ariblk.ttf"),
    "impact": ("impact.ttf", "impact.ttf"),
    "georgia": ("georgia.ttf", "georgiab.ttf"),
    "consolas": ("consola.ttf", "consolab.ttf"),
    "segoe ui": ("segoeui.ttf", "segoeuib.ttf"),
    "calibri": ("calibri.ttf", "calibrib.ttf"),
    "verdana": ("verdana.ttf", "verdanab.ttf"),
    "tahoma": ("tahoma.ttf", "tahomabd.ttf"),
    "times new roman": ("times.ttf", "timesbd.ttf"),
    "courier new": ("cour.ttf", "courbd.ttf"),
    "trebuchet ms": ("trebuc.ttf", "trebucbd.ttf"),
    # 本机常见但不在跨平台安全名单内的字体：给了度量映射，宽度才算得准，
    # 可移植性问题由下方 _SAFE_FONTS 单独告警，两件事分开报。
    "noto sans sc": ("NotoSansSC-VF.ttf", "NotoSansSC-VF.ttf"),
    "noto serif sc": ("NotoSerifSC-VF.ttf", "NotoSerifSC-VF.ttf"),
}
_REF = 1000  # 参考字号：度量后按比例换算，避免整数字号取整误差

# 跨平台「不会被静默替换」的字体白名单。只列 Windows 自带不够 ——
# deck 发给 Mac 用户时同样会被替换，所以含 macOS / 日韩常见系统字体。
_SAFE_FONTS = {
    "微软雅黑", "microsoft yahei", "simhei", "黑体", "simsun", "宋体",
    "kaiti", "楷体", "fangsong", "仿宋", "dengxian", "等线",
    "microsoft jhenghei", "pmingliu", "mingliu",
    "pingfang sc", "pingfang tc", "pingfang hk",
    "heiti sc", "heiti tc", "songti sc", "songti tc", "stsong",
    "yu gothic", "yu mincho", "meiryo", "ms gothic", "ms mincho",
    "malgun gothic", "gulim", "batang",
    "arial", "arial black", "calibri", "segoe ui", "verdana",
    "helvetica", "helvetica neue", "tahoma", "trebuchet ms",
    "times new roman", "times", "georgia", "cambria", "palatino",
    "garamond", "book antiqua", "consolas", "courier new", "menlo",
    "monaco", "impact",
}

_missing: set[str] = set()
_pil_missing = False
# 整类检查因环境缺失没跑：必须显式上报，绝不静默降级（否则「硬伤 0·告警 0」会被读成全部通过）。
_DEGRADED: list[str] = []

@lru_cache(maxsize=64)
def _face(family: str, bold: bool):
    """加载真实字体；缺失返回 None（调用方回落估算）。"""
    global _pil_missing
    try:
        from PIL import ImageFont
    except ImportError:
        _pil_missing = True
        return None
    entry = _FONT_FILES.get((family or "").strip().lower())
    if not entry:
        # "?" 是 qa 自己给「未显式指定字体」的占位符，不是真字体，不进缺失名单
        if family and family.strip() and family != "?":
            _missing.add(family)  # 空字体名是病态输入，不进缺失名单（回落估算即可）
        return None
    path = _FONT_DIR / (entry[1] if bold else entry[0])
    if not path.exists():
        if family and family.strip() and family != "?":
            _missing.add(family)
        return None
    try:
        return ImageFont.truetype(str(path), _REF)
    except Exception:
        if family and family.strip() and family != "?":
            _missing.add(family)
        return None

def _wide(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in ("W", "F")

def width_pt(text: str, family: str, bold: bool, pt: float) -> float:
    """文字在给定字体字号下的真实前进宽度（pt）。

    拉丁专用字体（Arial Black 等）没有 CJK 字形，中文渲染时会回退到默认中文字体，
    宽度必须按回退字体量——直接量原字体会得出近零宽度，超宽检查全部失明。
    """
    low = (family or "").strip().lower()
    if low not in _LATIN_ONLY:
        f = _face(family, bold)
        if f is not None:
            try:
                return f.getlength(text) / _REF * pt
            except Exception:
                pass
        return sum(pt * (1.0 if _wide(c) else 0.55) for c in text)
    total, i, n = 0.0, 0, len(text)
    while i < n:
        j = i + 1
        wide = _wide(text[i])
        while j < n and _wide(text[j]) == wide:
            j += 1
        seg = text[i:j]
        f = _face("微软雅黑" if wide else family, bold)
        done = False
        if f is not None:
            try:
                total += f.getlength(seg) / _REF * pt
                done = True
            except Exception:
                pass
        if not done:
            total += sum(pt * (1.0 if _wide(c) else 0.55) for c in seg)
        i = j
    return total

def wrapped_lines(text: str, box_w_in: float, family: str, bold: bool, pt: float) -> int:
    cap = box_w_in * 72.0
    n, cur = 1, 0.0
    for ch in text:
        if ch in ("\n", "\v"):  # \v 是 <a:br/> 软换行，真实断行，不是字符
            n, cur = n + 1, 0.0
            continue
        w = width_pt(ch, family, bold, pt)
        if cur + w > cap + 1e-6:
            n, cur = n + 1, w
        else:
            cur += w
    return n

# ---------------------------------------------------------------- 对比度

def _lum(rgb: int) -> float:
    def lin(v: float) -> float:
        v /= 255.0
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (rgb >> 16) & 255, (rgb >> 8) & 255, rgb & 255
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)

def contrast_ratio(a: int, b: int) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)

def _rgb_of(color) -> int | None:
    """MSO 颜色为显式 RGB 时返回 0xRRGGBB；主题色 / 未设色返回 None。"""
    try:
        from pptx.enum.dml import MSO_COLOR_TYPE
        if color is None or color.type != MSO_COLOR_TYPE.RGB:
            return None
        return int(str(color.rgb), 16)
    except Exception:
        return None

def _fill_of(sh):
    """形状的 FillFormat；没有填充属性的形状（图片、图表 GraphicFrame）返回 None。"""
    return sh.fill if hasattr(sh, "fill") else None

# 拉丁专用字体：没有 CJK 字形，中文排进去会静默回退成默认中文字体
_LATIN_ONLY = {
    "arial", "arial black", "impact", "georgia", "times new roman", "consolas",
    "segoe ui", "verdana", "tahoma", "trebuchet ms", "calibri", "courier new",
    "cambria", "palatino", "garamond", "book antiqua", "menlo", "monaco",
    "helvetica", "helvetica neue",
}

def _solid_rgb(fill) -> int | None:
    """填充为纯色时返回 0xRRGGBB，否则 None（含无填充 / 主题色）。"""
    try:
        from pptx.enum.dml import MSO_FILL
        if fill is None or fill.type != MSO_FILL.SOLID:
            return None
        return _rgb_of(fill.fore_color)
    except Exception:
        return None

def _img_avg_rgb(shape) -> int | None:
    """图片不透明区域的「墨色」= 最暗一档像素的平均色（0xRRGGBB）。

    用于图片型强调元素的对比度兜底。用最暗一档而不是全图均值：线稿插图（figure:*）
    大量像素是稀疏线条与浅色阴影，全图均值会被浅色区拉白、把深色线误判成"几乎不可见"；
    线与字形本身的颜色才是决定可读性的那个颜色。
    """
    try:
        import io
        from PIL import Image
        im = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA").resize((48, 48))
        px = [(r, g, b) for r, g, b, a in im.getdata() if a > 32]
        if not px:
            return None
        px.sort(key=lambda p: 0.299 * p[0] + 0.587 * p[1] + 0.114 * p[2])
        core = px[:max(1, len(px) // 4)]           # 最暗的 1/4
        n = len(core)
        return ((sum(p[0] for p in core) // n) << 16
                | (sum(p[1] for p in core) // n) << 8
                | (sum(p[2] for p in core) // n))
    except Exception:
        return None

def _translucent(shape) -> bool:
    """形状 spPr 里有 <60% 的 alpha 就当半透明处理（跳过对比度检查）。"""
    try:
        xml = shape._element.spPr.xml
    except Exception:
        return False
    m = re.search(r'<a:alpha val="(\d+)"', xml)
    return bool(m and int(m.group(1)) < 60000)

def _bg_under(shapes, idx: int, slide):
    """文字形状的等效底色：自身填充 → 下方最近的不透明纯色覆盖块 → 页面背景。"""
    sh = shapes[idx]
    own = _solid_rgb(_fill_of(sh))
    if own is not None and not _translucent(sh):
        return own
    try:
        cx = (sh.left + sh.width / 2) / EMU
        cy = (sh.top + sh.height / 2) / EMU
    except Exception:
        return None
    candidate = None
    for other in shapes[:idx]:
        if other.has_text_frame and other.text_frame.text.strip():
            continue  # 只拿非文字形状当底（文字形状互相叠是另一条检查）
        try:
            contains = (other.left / EMU <= cx <= (other.left + other.width) / EMU
                        and other.top / EMU <= cy <= (other.top + other.height) / EMU)
        except Exception:
            continue
        if not contains:
            continue
        # 压在图片/图表/表格上：真实底色未知（页面底色被盖住了），跳过对比度检查。
        # 若回退到页面底色，白字压深色照片会被误判成压浅底 → 硬伤误报。
        st_o = getattr(other, "shape_type", None)
        if st_o == 13 or getattr(other, "has_chart", False) or getattr(other, "has_table", False):
            return None
        rgb = _solid_rgb(_fill_of(other))
        if rgb is not None and not _translucent(other):
            candidate = rgb  # 不提前 return：取 z 序最上层的那个，全页底矩形在最底层
    if candidate is not None:
        return candidate
    try:
        bg = _solid_rgb(slide.background.fill)
        if bg is not None:
            return bg
    except Exception:
        pass
    return None

# ---------------------------------------------------------------- 中文排版

# 行首禁则：这些标点不能落在一行开头
_NO_START = set("，。、；：！？）》」』】〉”’%,.;:!?)]}")
# 行尾禁则：这些标点不能落在一行末尾
_NO_END = set("（《「『【〈“‘([{")

def _ideo(ch: str) -> bool:
    return "\u3400" <= ch <= "\u9fff" or "\uf900" <= ch <= "\ufaff"

try:
    import jieba
    jieba.setLogLevel(60)  # 压掉建词典时的 stderr 噪音
    jieba.initialize()  # 懒加载：不初始化 FREQ 永远是空词典，劈词检查会静默失效
    _WORDS = jieba.dt.FREQ
except ImportError:
    _WORDS = {}
    _DEGRADED.append("劈词：jieba 未安装（python -m pip install jieba）—— 跨行劈词检查整类关闭")
except Exception as _e:
    _WORDS = {}
    _DEGRADED.append(f"劈词：jieba 初始化失败（{str(_e)[:80]}）—— 跨行劈词检查整类关闭")

def _splits_word(prev: str, nxt: str) -> bool:
    """交界两字是否构成词典词（跨行劈开了一个词）。"""
    if not _WORDS or not prev or not nxt:
        return False
    return _WORDS.get(prev[-1] + nxt[0], 0) > 0

def _split_words_at(segs_a: list[str], segs_b: list[str], si: str, warn: list):
    """两段视觉行序列的交界处查劈词（段落间、<a:br/> 软换行处都算行界）。"""
    a = next((s for s in reversed(segs_a) if s), "")
    b = next((s for s in segs_b if s), "")
    if a and b and _ideo(a[-1]) and _ideo(b[0]) and _splits_word(a, b):
        warn.append(f"[{si}] 疑似劈词：…{a[-3:]} | {b[:2]}…")

# ---------------------------------------------------------------- 逐段排版度量

_LINE_MIN_MULT = 0.5

def _para_meta(para):
    runs = [r for r in para.runs if r.text.strip()]
    pt = max((r.font.size.pt for r in runs if r.font.size), default=None)
    fam = next((r.font.name for r in runs if r.font.name), None)
    bold = next((r.font.bold for r in runs if r.font.bold is not None), None)
    return runs, pt, fam, bold

def text_height_in(tf, w_eff_in: float, scale: float = 1.0) -> float:
    """逐段落累计真实行高：字号、行距、段前段后各段各算。

    scale 是 normAutofit fontScale 折算的缩字系数——渲染字号 = 标称 × scale，
    行高随之缩放；段前段后是绝对磅值，不随缩字变。
    """
    total_pt = 0.0
    for para in tf.paragraphs:
        _, pt, fam, bold = _para_meta(para)
        pt = (pt if pt is not None else 18.0) * scale
        fam = fam or "?"
        mult = 1.32
        ls = para.line_spacing
        if ls is not None:
            if hasattr(ls, "pt"):                     # Length：固定磅值行距
                mult = max(ls.pt / pt, _LINE_MIN_MULT)
            else:                                     # float：倍数行距
                mult = max(float(ls) * 1.32, _LINE_MIN_MULT)
        n = wrapped_lines(para.text, w_eff_in, fam, bool(bold), pt)
        total_pt += n * pt * mult
        for sp in (para.space_before, para.space_after):
            if sp is not None:
                total_pt += sp.pt
    return total_pt / 72.0

def _inset_in(v, default_in: float) -> float:
    return v / EMU if v is not None else default_in

# ---------------------------------------------------------------- 文字类检查（文本框 / 表格单元格共用）

def _autofit_info(tf):
    """读 bodyPr 的自适应配置，返回 (kind, scale_pct)。

    kind: None / "norm"（normAutofit，缩字适配）/ "sp"（spAutoFit，形状随文字长高）。
    spAutoFit 以存储尺寸渲染且 python-pptx 新建文本框模板默认带它——属库痕迹而非
    模型意图，静态度量仍有效，不当告警。normAutofit 的 fontScale 才真正改变渲染
    字号（实际字号 = 标称 × fontScale/100000），必须按实际值度量。
    """
    try:
        bxml = tf._txBody.bodyPr.xml
    except Exception:
        return None, None
    if "normAutofit" in bxml:
        m = re.search(r'normAutofit[^>]*fontScale="(\d+)"', bxml)
        return "norm", (int(m.group(1)) / 1000.0 if m else None)
    if "spAutoFit" in bxml:
        return "sp", None
    return None, None

def audit_text_frame(si, x, y, w, h, tf, bg_rgb, hard: list, warn: list, faces: set,
                     insets=None, scale: float = 1.0, small_agg: list | None = None,
                     rotated: bool = False):
    """对一个文本框架做全部文字类检查，返回字数当量（无可检文本返回 None）。

    bg_rgb 是调用方解析好的等效底色（None 跳过对比度类检查）；
    insets 为 (左,右,上,下) 英寸，表格单元格传入自身的边距，形状路径用 tf 默认；
    scale 是 normAutofit 缩字系数，渲染字号 = 标称字号 × scale；
    small_agg 传入列表时，小字硬伤聚合进列表（表格逐格报会刷屏），由调用方汇总；
    rotated=True（框旋转 ≥3°）时跳过溢出/行超宽/孤行这些依赖未旋转几何的检查
    ——旋转后有效宽高换轴，按框尺寸判必然误报（实测：竖排编号被判 147% 超宽），
    旋转本身另有提醒，视觉核对交闸门二。
    """
    runs = [r for p in tf.paragraphs for r in p.runs if r.text.strip()]
    if not runs:
        return None

    # 字号下限：逐 run 查，混合字号的框里藏小字也能抓到；有缩字按实际字号判
    small_runs = []
    for r in runs:
        if r.font.size:
            eff = r.font.size.pt * scale
            if eff < MIN_PT:
                if small_agg is None:
                    tag = (f"{r.font.size.pt:g}pt" if scale == 1.0
                           else f"{r.font.size.pt:g}pt×缩字{scale:.0%}={eff:.1f}pt")
                    hard.append(f"[S{si}] 字号 {tag} < {MIN_PT:g}pt | {r.text[:16]!r}")
                else:
                    small_runs.append(eff)
    if small_runs:
        small_agg.append((min(small_runs), len(small_runs)))

    fam0 = next((r.font.name for r in runs if r.font.name), "?")
    faces.add(fam0)

    # 中文排进拉丁专用字体：渲染静默回退成默认中文字体（衬线感的来源）
    bad_lat = sorted({r.font.name for r in runs
                      if r.font.name and r.font.name.strip().lower() in _LATIN_ONLY
                      and any(_wide(c) for c in r.text)})
    if bad_lat:
        warn.append(f"[S{si}] 中文用了拉丁专用字体 {bad_lat} —— 渲染会回退成默认中文字体，"
                    "字形与字宽都失控；中文用中文字体，拉丁/编号才用 Arial Black 一类")

    if insets is None:
        ml = _inset_in(tf.margin_left, 0.1)
        mr = _inset_in(tf.margin_right, 0.1)
        mt = _inset_in(tf.margin_top, 0.05)
        mb = _inset_in(tf.margin_bottom, 0.05)
    else:
        ml, mr, mt, mb = insets
    w_eff = max(w - ml - mr, 0.2)
    h_eff = h - mt - mb

    full = tf.text
    char_load = sum(1.0 if _wide(c) else (0.5 if c.isalnum() else 0.0) for c in full)

    # emoji 当图标：反 AI 清单明令禁止（用图标库或删除）
    emo = _EMOJI.search(full)
    if emo:
        warn.append(f"[S{si}] 文本里有 emoji {emo.group(0)!r} —— 反 AI 清单禁 emoji 当图标，"
                    "需要图标就用 inject_icons.py（Tabler 线性）")

    # 正文里的几何/数学符号：中文字体把它们画成汉字的四成高、笔画细一档，一眼假（pitfalls #29）
    glyphs = sorted(set(_SHAPE_GLYPH.findall(full)))
    if glyphs:
        warn.append(f"[S{si}] 正文里有排版级几何符号 {' '.join(glyphs)} —— 中文字体里这类符号"
                    "只有汉字四成大小、笔画细一档，看着就是 LLM 文本味：句子里的「三角形」直接写汉字，"
                    "数学式改成 $$…$$ 独占文本框走 inject_math，图形用 inject_figure/inject_icons（pitfalls #29）")

    need = text_height_in(tf, w_eff, scale)
    if not rotated and need > h_eff + 0.02:  # h_eff 为负（框比内边距还矮）直接判溢出，不跳过
        hard.append(f"[S{si}] 溢出 需 {need:.2f}\" 实有 {max(h_eff, 0):.2f}\" "
                    f"| {fam0} | {full[:26]!r}")

    # 对比度 + 未显式设色兜底
    fg = next((rgb for r in runs if (rgb := _rgb_of(r.font.color)) is not None), None)
    if bg_rgb is not None:
        pt0 = next((r.font.size.pt for r in runs if r.font.size), 18.0) * scale
        bold0 = bool(next((r.font.bold for r in runs if r.font.bold is not None), False))
        large = pt0 >= 18 or (pt0 >= 14 and bold0)
        if fg is not None:
            ratio = contrast_ratio(fg, bg_rgb)
            if ratio < 2.5:
                hard.append(f"[S{si}] 对比度 {ratio:.1f}:1 （#{fg:06X} 压 #{bg_rgb:06X}） "
                            f"几乎不可读 | {full[:16]!r}")
            elif ratio < (3.0 if large else 4.5):
                warn.append(f"[S{si}] 对比度 {ratio:.1f}:1 （#{fg:06X} 压 #{bg_rgb:06X}，"
                            f"{pt0:g}pt {'粗体' if bold0 else '常规'}）低于 {'3.0' if large else '4.5'} "
                            f"| {full[:16]!r}")
        elif (all(_rgb_of(r.font.color) is None for r in runs)
              and _lum(bg_rgb) < 0.3):
            warn.append(f"[S{si}] 文字无显式 RGB（继承主题或用了主题色）压深底 "
                        f"#{bg_rgb:06X} —— 对比度无法验证，渲染若成空白即此因")

    cap = w_eff * 72.0
    paras = list(tf.paragraphs)
    prev_segs: list[str] = []
    for i, para in enumerate(paras):
        raw = para.text
        if not raw.strip():
            continue
        _, pt_p, fam_p, bold_p = _para_meta(para)
        pt_p = pt_p if pt_p is not None else 18.0
        fam_p = fam_p or "?"
        segs = raw.split("\v")  # <a:br/> 软换行是真实行界，逐段检查
        for seg in segs:
            if not seg:
                continue
            used = width_pt(seg, fam_p, bool(bold_p), pt_p * scale)
            if not rotated and used > cap + 1e-6:
                hard.append(f"[S{si}] 行超宽 {used/cap*100:.0f}% → 会二次折行 "
                            f"| {pt_p * scale:g}pt {fam_p} | {seg[:22]!r}")
            if seg[0] in _NO_START:
                hard.append(f"[S{si}] 行首禁则：以 {seg[0]!r} 开头 | {seg[:18]!r}")
            if seg[-1] in _NO_END:
                hard.append(f"[S{si}] 行尾禁则：以 {seg[-1]!r} 结尾 | {seg[:18]!r}")
        _split_words_at(prev_segs, segs, f"S{si}", warn)          # 上一段落 ↔ 本段
        for a, b in zip(segs, segs[1:]):                          # 段内软换行处
            _split_words_at([a], [b], f"S{si}", warn)
        if not rotated and len(paras) >= 2 and i == len(paras) - 1 and segs and segs[-1]:
            last = segs[-1]
            used_last = width_pt(last, fam_p, bool(bold_p), pt_p * scale)
            if 0 < used_last < cap * 0.30:
                warn.append(f"[S{si}] 孤行尾行（{used_last/cap*100:.0f}% 宽）| {last[:18]!r}")
        prev_segs = segs

    return char_load

# ---------------------------------------------------------------- 交付合法性

def _brief_names(names, cap: int = 6) -> str:
    """把形状名集合缩成一行可读清单（过长则截断并报总数）——morph 配对诊断用。"""
    ns = sorted(n for n in names if n)
    if not ns:
        return "（无）"
    head = "、".join(ns[:cap])
    return head + ("…（共 %d 个）" % len(ns) if len(ns) > cap else "")


# 真正产生"运动"的行为元素；可见性设置（p:set，注入器固定写 dur=500）不算时长，
# 否则它会把短动画盖住（实测 0.1s 的淡入被算成 0.5s，过短告警永远不响）。
_MOTION_TAGS = ("animEffect", "anim", "animScale", "animRot", "animMotion")


def _walk_eff(node, base, items):
    """childTnLst 里递归累加 delay，收集 (组内相对起始ms, 时长ms)。"""
    for par in node:
        if par.tag != f"{{{P_NS}}}par":
            continue
        ctn = par.find(f"{{{P_NS}}}cTn")
        if ctn is None:
            continue
        delay = 0
        for cond in ctn.findall(f"{{{P_NS}}}stCondLst/{{{P_NS}}}cond"):
            d = (cond.get("delay") or "").strip()
            if d.isdigit():
                delay += int(d)
        start = base + delay
        if ctn.get("nodeType") in ("clickEffect", "withEffect", "afterEffect"):
            durs = []
            for beh in ctn.iter():
                if beh.tag.rsplit("}", 1)[-1] not in _MOTION_TAGS:
                    continue
                durs += [int(x.get("dur")) for x in beh.iter(f"{{{P_NS}}}cTn")
                         if (x.get("dur") or "").isdigit()]
            items.append((start, max(durs) if durs else 0))
        inner = ctn.find(f"{{{P_NS}}}childTnLst")
        if inner is not None:
            _walk_eff(inner, start, items)
    return items


def _anim_groups(timing):
    """timing -> [(是否点击组, [(组内相对起始ms, 时长ms), …]), …]

    组 = mainSeq 子层每一个 p:par，一组对应一次点击或一次自动起播（"一次揭示出几个"
    就是"这一组里有几条"）。组内起始时刻沿路径累加 delay 求相对值——**不逐字比对
    delay 属性**，因为 PowerPoint 会重写嵌套编码。
    """
    seq = timing.find(f".//{{{P_NS}}}cTn[@nodeType='mainSeq']")
    if seq is None:
        return []
    ctl = seq.find(f"{{{P_NS}}}childTnLst")
    if ctl is None:
        return []
    out = []
    for gpar in ctl:
        if gpar.tag != f"{{{P_NS}}}par":
            continue
        gctn = gpar.find(f"{{{P_NS}}}cTn")
        click = gctn is not None and any(
            (c.get("delay") or "") == "indefinite"
            for c in gctn.iter(f"{{{P_NS}}}cond"))
        inner = gctn.find(f"{{{P_NS}}}childTnLst") if gctn is not None else None
        out.append((click, _walk_eff(inner, 0, []) if inner is not None else []))
    return out


def check_delivery(path: str) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    try:
        zf = zipfile.ZipFile(path)
    except Exception as e:
        return [("硬伤", f"不是合法 zip：{e}")]
    with zf:
        names = set(zf.namelist())
        if "[Content_Types].xml" not in names:
            out.append(("硬伤", "缺少 [Content_Types].xml，PowerPoint 会拒绝打开"))
        broken = zf.testzip()
        if broken:
            out.append(("硬伤", f"压缩包成员损坏：{broken}"))

        # 放映设置：计时自动翻页 / 展台循环——整份会自己翻，演讲时按不住
        if "ppt/presentation.xml" in names:
            try:
                pxml = zf.read("ppt/presentation.xml").decode("utf-8", errors="replace")
            except Exception:
                pxml = ""
            sp = re.search(r"<p:showPr\b[^>]*>", pxml)
            if sp:
                kv = dict(re.findall(r'([\w:]+)="([^"]*)"', sp.group(0)))
                if kv.get("useTimings") in ("1", "true"):
                    out.append(("告警", "[deck] 演示文稿开了「使用计时/排练计时」"
                                        "（p:showPr@useTimings）—— 放映时按计时自动翻页，"
                                        "讲到一半就跑掉；在「幻灯片放映 → 设置放映方式」里关掉（若就是要按计时自动播放，写明即可）"))
                if kv.get("showType") == "kiosk":
                    out.append(("告警", "[deck] 放映方式是「展台(全屏幕)循环」"
                                        "（p:showPr@showType=\"kiosk\"）—— 会自动循环、无法"
                                        "人工控制节奏；演讲场合改成「演讲者放映」（若就是要无人值守循环，那它是要的功能，写明即可）"))

        prev_names: set[str] = set()
        for slide in sorted(n for n in names
                            if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)):
            blob = zf.read(slide)
            try:
                root = _xml_parse(blob)
            except Exception as e:
                out.append(("硬伤", f"{slide} XML 解析失败：{e}"))
                continue
            # 单段落多个 <a:pPr>：违反 schema，只有 PowerPoint 会重排错乱
            for p in root.iter(f"{{{A_NS}}}p"):
                if sum(1 for c in p if c.tag == f"{{{A_NS}}}pPr") > 1:
                    out.append(("硬伤", f"{slide} 有段落含多个 <a:pPr>，"
                                        "PowerPoint 打开后该行会错乱（pitfalls #4）"))
                    break
            blurs = {int(m) for m in re.findall(rb'blurRad="(\d+)"', blob)}
            if blurs:
                soft = {b for b in blurs if b > 12700}
                if soft and len(soft) < len(blurs):
                    # 逐页取 min 会互相掩护：同页一硬多柔时硬伤被藏起来（实战踩过）
                    out.append(("告警", f"{slide} 同页混有硬阴影与柔和阴影（blurRad {min(blurs)} vs "
                                        f"{max(soft) / 12700:.0f}pt）—— 若风格要求统一零模糊硬阴影"
                                        "则已部分失效（pitfalls #2）"))
                elif soft:
                    out.append(("告警", f"{slide} 阴影 blurRad 最小 {min(blurs) / 12700:.0f}pt，"
                                        "若风格要求零模糊硬阴影则已失效（pitfalls #2）"))
            hash_colors = re.findall(rb'val="#[0-9A-Fa-f]{6}"', blob)
            if hash_colors:
                out.append(("硬伤", f"{slide} 颜色值带 #（{len(hash_colors)} 处）—— 轻则丢样式重则"
                                    "文件损坏，色值只写六位十六进制（pitfalls #7）"))

            # ---- 公式残留：未转换的 LaTeX 会显示为乱码文本（inject_math.py 漏网）----
            xml_txt = blob.decode("utf-8", errors="replace")
            for dollar in re.findall(r"\$\$([^$\n]{1,80})\$\$|\$([^$\n]{1,60})\$", xml_txt):
                cand = dollar[0] or dollar[1]
                if re.search(r"[\\^_]", cand):     # 有 LaTeX 特征才算公式；$5、$8 是价格
                    out.append(("硬伤", f"{slide} 公式未转换（会显示为乱码文本）：${cand[:30]}$ "
                                        "—— 先跑 inject_math.py"))
                    break

            # ---- 动画纪律（inject_anim.py 注入的 timing / transition）----
            si_no = int(re.search(r"(\d+)", slide).group(1))
            shape_names = {e.get("name") for e in root.iter(f"{{{P_NS}}}cNvPr")
                           if e.get("name")}
            shape_ids = {int(e.get("id")) for e in root.iter(f"{{{P_NS}}}cNvPr")
                         if e.get("id") and (e.get("id") or "").isdigit()}
            spids = {int(t.get("spid")) for t in root.iter(f"{{{P_NS}}}spTgt")
                     if t.get("spid") and (t.get("spid") or "").isdigit()}
            dangling = spids - shape_ids
            if dangling:
                out.append(("硬伤", f"{slide} 动画引用了不存在的形状 id {sorted(dangling)}"
                                    " —— 打开会提示修复或动画被静默丢弃"))
            timing = root.find(f"{{{P_NS}}}timing")
            if timing is not None:
                effects = [c for c in timing.iter(f"{{{P_NS}}}cTn")
                           if c.get("nodeType") in ("clickEffect", "withEffect", "afterEffect")]
                # ---- 节拍：一次揭示出几个、共几组、每组播多久 ----
                # 组 = 一次点击或一次自动起播；「一次点击出几个」就是「这组里有几条」。
                grp = _anim_groups(timing)
                n_groups = len(grp)
                clicks = sum(1 for c, _ in grp if c)
                eff_durs = [d for _c, gitems in grp for _s, d in gitems]
                if n_groups > 4:
                    out.append(("告警", f"{slide} 一页有 {n_groups} 个揭示节拍（其中点击 {clicks} 次）"
                                        " —— 演讲时手上太忙；把同一论点里的东西并成一组，"
                                        "一页控制在 4 组以内（炫技/展映档可明知而保留——规格书与交付说明写明即可）"))
                singles = [1 for _c, gitems in grp if len(gitems) == 1]
                if n_groups >= 3 and len(singles) >= 2:
                    out.append(("告警", f"{slide} 有 {len(singles)} 个节拍只含一个对象（共 {n_groups} 组）"
                                        " —— 节拍过碎、等于一次点一下：同一句话里的东西应该同一次"
                                        "出现（一组里的后续条目写「之后」或「同时」）。炫技/展映档可明知而保留"))
                for gi, (_c, gitems) in enumerate(grp, 1):
                    if not gitems:
                        continue
                    span = max(s + d for s, d in gitems) - min(s for s, _ in gitems)
                    if span > 8000:
                        out.append(("告警", f"{slide} 第 {gi} 个节拍要播 {span / 1000:.1f}s —— 这一次"
                                            "点击之后观众得干等；拆成两组，或缩短单条时长。炫技/展映档可明知而保留"))
                short = [d for d in eff_durs if 0 < d < 200]
                if short:
                    out.append(("告警", f"{slide} 有 {len(short)} 条动画短于 0.2s（最短 {min(short)}ms）"
                                        " —— 一闪而过，观众看不见；入场至少 0.3s（刻意做的瞬间闪现除外，须在规格书写明）"))
                if len(effects) > 8:
                    out.append(("告警", f"{slide} {len(effects)} 条对象动画 —— 逐条登场会拖垮节奏，"
                                        "每页入场预算 ≤8 条；超了就该拆页，或砍掉装饰性的那几条（炫技/展映档可明知而保留）"))
                slow = [d for d in eff_durs if d > 3000]
                if slow:
                    out.append(("告警", f"{slide} 有 {len(slow)} 条动画长于 3s（最长 "
                                        f"{max(slow) / 1000:.1f}s）—— 拖沓；单条建议 ≤2s，"
                                        "要显示得久就让下一条「之后」接上"))
            # ---- 自动换片时间：这页会自己翻过去（讲到一半就跑掉）----
            for tr in root.iter(f"{{{P_NS}}}transition"):
                adv_tm = tr.get("advTm")
                if adv_tm or tr.get("advClick") == "0":
                    out.append(("告警", f"{slide} 带自动换片时间或禁止点击换页（"
                                        f"advTm={adv_tm or '—'} / advClick={tr.get('advClick', '—')}）"
                                        " —— 放映时这页会自己翻过去、演讲按不住。注意 p14:dur 是"
                                        "切换本身的播放时长，不是换片计时；这项多半来自模板或别的"
                                        "工具，去掉它（展映/无人值守/大屏循环档下「自动翻页」正是"
                                        "想要的功能——须是用户有意设置，并在规格书与交付说明里写明）"))
                    break
            if b":morph" in blob:
                shared = shape_names & prev_names
                if not shared:
                    out.append(("告警", f"{slide} 用了平滑(morph)切换，但和上一页没有一个同名形状"
                                        f" —— 会整页退化成普通淡入。本页独有：{_brief_names(shape_names)}；"
                                        f"上一页独有：{_brief_names(prev_names)}。"
                                        "要连续变形的元素（同一母题块、同一张图）在相邻两页用同一个 "
                                        "objectName"))
            prev_names = shape_names

        for chart in sorted(n for n in names
                            if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)):
            blob = zf.read(chart)
            if (re.search(rb'val="(?:percentStacked|stacked)"', blob)
                    and b'dLblPos val="outEnd"' in blob):
                out.append(("硬伤", f"{chart} 堆叠柱状图用 outEnd 数据标签，"
                                    "PowerPoint 会提示修复并丢弃图表（pitfalls #11）"))
            hash_colors = re.findall(rb'val="#[0-9A-Fa-f]{6}"', blob)
            if hash_colors:
                out.append(("硬伤", f"{chart} 颜色值带 #（{len(hash_colors)} 处）—— "
                                    "色值只写六位十六进制（pitfalls #7）"))
            for ptag in (b"pieChart", b"doughnutChart"):
                m = re.search(rb"<c:" + ptag + rb"[ >].*?</c:" + ptag + rb">", blob, re.S)
                if not m:
                    continue
                mv = re.search(rb"<c:val>.*?</c:val>", m.group(0), re.S)
                npts = len(re.findall(rb"<c:pt idx", mv.group(0))) if mv else 0
                if npts > 5:
                    out.append(("告警", f"{chart} 饼图 {npts} 块 > 5 —— 小项合并或改横向条形图，"
                                        "两块 <6% 的扇区肉眼难分（styles.md 图表纪律）"))

    try:
        prs = Presentation(path)
        for i, sl in enumerate(prs.slides, 1):
            for rel in sl.part.rels.values():
                if rel.is_external:
                    continue
                try:
                    if not rel.target_part.blob:
                        out.append(("硬伤", f"S{i} 关系 {rel.rId} 指向空 part，图片会是空框"))
                except Exception as e:
                    out.append(("硬伤", f"S{i} 关系 {rel.rId} 无法解析：{e}"))
    except Exception as e:
        out.append(("硬伤", f"python-pptx 无法完整读取：{e}"))
    return out

# --------------------------------------------- 报告排版（分诊：摘要前置、按页分组）

_PAGE_TAG = re.compile(r"^\[S(\d+)\]")

def _group_key(line: str) -> tuple[int, int]:
    """分组键：带 [S页码] 的行按页号升序，全局行（[交付]/[字体]/[画布]/[deck]）排最后。"""
    m = _PAGE_TAG.match(line)
    return (0, int(m.group(1))) if m else (1, 0)

def _ordered(lines: list[str]) -> list[str]:
    return sorted(lines, key=_group_key)  # 稳定排序：组内保持插入序

def _report_lines(hard, warn, degraded, n_texts: int, metrics: str) -> list[str]:
    """排成一屏可读的报告：摘要 → 计数 → 降级块 → 硬伤（按页） → 告警（按页）。

    分诊靠顺序与分组，不做截断——藏掉真实问题比报告长更糟。
    """
    out: list[str] = []
    head = f"\n{n_texts} 个文本箱体 · 度量：{metrics}"
    if degraded:
        head += f" · 检查降级 {len(degraded)}"
    out.append(head)
    count = f"\n硬伤 {len(hard)} · 告警 {len(warn)}"
    if degraded:
        count += f" · 检查降级 {len(degraded)}"
    out.append(count)
    if degraded:
        out.append("  ⚠ 检查降级（下列检查整类没跑，不代表通过）：")
        for d in degraded:
            out.append(f"    [降级] {d}")
    for marker, items in (("✗", hard), ("?", warn)):
        ordered = _ordered(items)
        multi = len({_group_key(ln) for ln in ordered if _group_key(ln)[0] == 0}) > 1
        cur = None
        for ln in ordered:
            k = _group_key(ln)
            if multi and k[0] == 0 and k != cur:
                out.append(f"  ── 第 {k[1]} 页 ──")  # 单页报告不插无谓分隔行
            out.append(f"  {marker} {ln}")
            cur = k
    return out

# ---------------------------------------------------------------- 主流程

def main() -> int:
    argv = sys.argv[1:]
    known = ("--min-pt", "--max-chars", "--palette")
    for a in argv:
        if a.startswith("--"):
            base = a.split("=", 1)[0]
            if base not in known:
                sys.exit(f"未知参数 {base}；可用：{' '.join(known)}")
    skip, args = False, []
    for a in argv:
        if skip:
            skip = False
            continue
        if a.startswith("--"):
            if "=" not in a and a in ("--min-pt", "--max-chars"):
                skip = True  # 「--flag 值」形式：下一个 token 是值，不是文件名
            continue
        args.append(a)
    name = args[0] if args else "deck.pptx"
    try:
        prs = Presentation(name)
    except Exception as e:
        sys.exit(f"打不开 {name}：{e}")
    SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
    print(f"{name} | canvas {SW:.2f} x {SH:.2f} in | min {MIN_PT:g}pt · max {MAX_CHARS} 字当量/页")

    hard: list[str] = []
    warn: list[str] = []
    texts: list[tuple] = []
    faces: set[str] = set()
    header_fp: dict[int, tuple | None] = {}
    band_pages: dict[int, bool] = {}
    bare_pages: dict[int, bool] = {}

    # 画布检查：pptxgenjs 默认 10×5.625（pitfalls #9），忘设 LAYOUT_WIDE 必系统性溢出
    if SW < 12:
        warn.append(f"[画布] 宽 {SW:.2f} 英寸是 pptxgenjs 默认/LAYOUT_4x3 —— 全部字号建议按 "
                    "13.33×7.5 调，不设 LAYOUT_WIDE 必系统性溢出（pitfalls #9）；"
                    "刻意用小画布则字号与边距 ×0.75")

    for si, slide in enumerate(prs.slides, 1):
        try:
            shapes = list(slide.shapes)
        except Exception as e:
            # 单页 XML 坏掉不该毁掉整份报告：记硬伤、跳过该页、其余页照常查
            hard.append(f"[S{si}] 页面结构无法解析（{type(e).__name__}：{e}）—— 该页跳过检查")
            header_fp[si] = None
            band_pages[si] = False
            bare_pages[si] = True
            continue
        fp, band, carriers = None, False, 0
        char_load = 0.0
        rotated = 0.0

        for idx, sh in enumerate(shapes):
            x, y = sh.left / EMU, sh.top / EMU
            w, h = sh.width / EMU, sh.height / EMU
            st = getattr(sh, "shape_type", None)
            rot = abs(getattr(sh, "rotation", 0) or 0)
            if rot > rotated:
                rotated = rot
            # 全宽色带：页脚（y>85%）或页眉（y<12%）的全宽窄条，有不透明纯色填充。
            # addText + fill 画的色带是 TEXT_BOX，所以判据看填充、不看形状类型；全页背景不算。
            if (w > SW - 1.5 and h < SH * 0.25 and (y > SH * 0.85 or y < SH * 0.12)
                    and _solid_rgb(_fill_of(sh)) is not None):
                band = True
            if st == 13 or getattr(sh, "has_chart", False):
                carriers += 1
            if st == 13:
                # 图片拉伸变形：按裁切（cover）折算后的原图比例 vs 放置框比例，差 >15% 即变形
                try:
                    native_w, native_h = sh.image.size
                    src_w = native_w * (1 - sh.crop_left - sh.crop_right)
                    src_h = native_h * (1 - sh.crop_top - sh.crop_bottom)
                    if src_w > 1 and src_h > 1 and h > 0:
                        src_ar, box_ar = src_w / src_h, w / h
                        if max(box_ar / src_ar, src_ar / box_ar) > 1.15:
                            warn.append(f"[S{si}] 图片拉伸变形：原图比例 {src_ar:.2f} 被拉成 "
                                        f"{box_ar:.2f} —— 比例不符就裁切，不要拉伸（pitfalls #14）")
                        # 有效 PPI：裁切后像素 ÷ 放置英寸。<72 必糊；<150 投屏发虚
                        if w > 0.05 and h > 0.05:
                            ppi = min(src_w / w, src_h / h)
                            if ppi < 72:
                                hard.append(f"[S{si}] 图片有效 PPI 仅 {ppi:.0f}"
                                            f"（{int(src_w)}px 放在 {w:.2f}\" 宽）—— 必糊："
                                            "换更大原图、缩小放置，或用 prep_assets.py 按目标尺寸出图")
                            elif ppi < 150:
                                warn.append(f"[S{si}] 图片有效 PPI {ppi:.0f} 偏低 —— 投屏会发虚，"
                                            "建议 ≥300（inject_math/inject_icons/inject_figure 产出的图天然 600–800）")
                    # 图片型强调元素（icon:*/formula:*/figure:*）对比度兜底：图片不进文字对比度检查，
                    # 深底上用暗色强调（如红公式压蓝图蓝）会静默不可读（pitfalls #24）
                    nm = getattr(sh, "name", "") or ""
                    if nm.startswith(("icon:", "formula:", "figure:")):
                        bgp = _bg_under(shapes, idx, slide)
                        avg = _img_avg_rgb(sh)
                        if bgp is not None and avg is not None:
                            ratio = contrast_ratio(avg, bgp)
                            kind = nm.split(":")[0]
                            if ratio < 2.5:
                                hard.append(f"[S{si}] {kind} 图形对比度 {ratio:.1f}:1"
                                            f"（#{avg:06X} 压 #{bgp:06X}）—— 几乎不可见："
                                            "深底上用亮色，或给强调元素加浅底托板（pitfalls #24）")
                            elif ratio < 3.0:
                                warn.append(f"[S{si}] {kind} 图形对比度 {ratio:.1f}:1 偏低 —— "
                                            "投影会发虚（pitfalls #24）")
                except Exception:
                    pass
            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                hard.append(f"[S{si}] 出画布 @({x:.2f},{y:.2f},{w:.2f}×{h:.2f})")

            # 表格：逐单元格走同一条文字检查路径；不进重叠检查（相邻 cell 共享边框）。
            # _Cell 没有位置属性（python-pptx 1.0.2），从表框原点 + 列宽/行高累加推算。
            if getattr(sh, "has_table", False):
                tbl = sh.table
                col_w = [(c.width or 914400) / EMU for c in tbl.columns]
                row_h = [(r.height or 914400) / EMU for r in tbl.rows]
                col_x = [0.0]
                for cw_ in col_w:
                    col_x.append(col_x[-1] + cw_)
                row_y = [0.0]
                for rh_ in row_h:
                    row_y.append(row_y[-1] + rh_)

                # 表格外框 vs 内容尺寸。pptxgenjs 的 addTable 只回填行高、不回填外框：
                # 缺 h 时外框高度退化成硬编码 1 英寸（源码 `cy || EMU`），缺 w 时宽退化成
                # 75% 页宽。真 PowerPoint 按外框裁剪，WPS/LibreOffice 自动撑开 → 渲染预览
                # 100% 看不见（pitfalls #30）。用原始 a:tr@h / a:gridCol@w 核对。
                r_raw = [r.height for r in tbl.rows]
                c_raw = [c.width for c in tbl.columns]
                if c_raw and all(v is not None for v in c_raw) and sum(c_raw) / EMU > w + 0.02:
                    hard.append(f"[S{si}] 表格列宽合计 {sum(c_raw) / EMU:.2f}\" > 外框宽 {w:.2f}\" —— "
                                "PowerPoint 会横向裁掉/溢出：给 w 且让 colW 之和 = w（pitfalls #30）")
                if r_raw and all(v is not None for v in r_raw) and sum(r_raw) > 0:
                    row_sum = sum(r_raw) / EMU
                    if h < row_sum - 0.02:
                        hard.append(f"[S{si}] 表格外框高 {h:.2f}\" < 行高合计 {row_sum:.2f}\" —— "
                                    "PowerPoint 会裁掉超出部分（WPS/LibreOffice 会自动撑开，预览看不出来）："
                                    "addTable 的 h 必须 = rowH × 行数（pitfalls #30）")
                    elif h > row_sum + 0.15:
                        warn.append(f"[S{si}] 表格外框高 {h:.2f}\" 比行高合计 {row_sum:.2f}\" 高 "
                                    f"{h - row_sum:.2f}\" —— 底部会留白")
                else:
                    warn.append(f"[S{si}] 表格行高未显式给出或为 0（靠 PowerPoint 自动定高）—— "
                                "外框高度与裁剪无法核对：给 rowH，并 let h = rowH × 行数（pitfalls #30）")
                small_agg: list = []
                for ri, row in enumerate(tbl.rows):
                    for ci, cell in enumerate(row.cells):
                        if cell.is_spanned or ci >= len(col_w) or ri >= len(row_h):
                            continue
                        cw = col_w[ci]
                        chh = row_h[ri]
                        if cell.is_merge_origin:
                            cw = sum(col_w[ci:ci + cell.span_width])
                            chh = sum(row_h[ri:ri + cell.span_height])
                        cbg = _solid_rgb(cell.fill)
                        if cbg is None:
                            cbg = _bg_under(shapes, idx, slide)
                        insets = tuple(
                            (m / EMU if m is not None else d)
                            for m, d in zip((cell.margin_left, cell.margin_right,
                                             cell.margin_top, cell.margin_bottom),
                                            (0.1, 0.1, 0.05, 0.05)))
                        load = audit_text_frame(si, x + col_x[ci], y + row_y[ri],
                                                cw, chh, cell.text_frame,
                                                cbg, hard, warn, faces, insets,
                                                small_agg=small_agg)
                        if load:
                            char_load += load
                if small_agg:  # 整表按字号分组汇总，逐格报会刷屏
                    by_pt: dict[int, int] = {}
                    for eff_pt, cnt in small_agg:
                        key = round(eff_pt, 1)
                        by_pt[key] = by_pt.get(key, 0) + cnt
                    for pt_, cnt in sorted(by_pt.items()):
                        hard.append(f"[S{si}] 表格 {cnt} 个单元格字号 {pt_:g}pt < {MIN_PT:g}pt")
                continue

            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            tf = sh.text_frame
            if fp is None and y < 1.2 and x < 3.0:
                fp = (round(x, 2), round(y, 2), round(w, 2), round(h, 2))
            # 自动缩字：normAutofit 的 fontScale 真正改变渲染字号，按实际值度量；
            # 未记录缩放比时无法折算，只能提醒交给闸门二。spAutoFit 静默（见 _autofit_info）。
            kind, scale_pct = _autofit_info(tf)
            if kind == "norm" and scale_pct is None:
                warn.append(f"[S{si}] 文本框配置了自动缩字（未记录缩放比）—— 渲染字号可能与标称"
                            f"不一致，闸门二核对 | {tf.text[:16]!r}")
            eff_scale = scale_pct / 100.0 if (kind == "norm" and scale_pct) else 1.0
            bg = _bg_under(shapes, idx, slide)
            load = audit_text_frame(si, x, y, w, h, tf, bg, hard, warn, faces,
                                    scale=eff_scale, rotated=rot >= 3)
            if load:
                char_load += load
            texts.append((x, y, w, h, tf.text[:18], si))

        if rotated >= 3:
            warn.append(f"[S{si}] 页内有旋转 {rotated:.0f}° 的形状 —— qa 的框检查不含旋转扫过范围"
                        "（pitfalls #16），闸门二必须核对渲染图")
        icon_count = sum(1 for sh in shapes
                         if (getattr(sh, "name", "") or "").startswith("icon:"))
        if icon_count > 3:
            warn.append(f"[S{si}] 本页 {icon_count} 个图标 —— 图标是路标不是装饰，每页 ≤3；"
                        "规格书应声明图标政策（家族/颜色角色/落点）")
        if char_load > MAX_CHARS:
            warn.append(f"[S{si}] 全页 {char_load:g} 字当量 > {MAX_CHARS} —— "
                        "信息密度偏高，投屏页考虑拆页或删减")
        header_fp[si] = fp
        band_pages[si] = band
        bare_pages[si] = carriers == 0

    for a, b in combinations(texts, 2):
        if a[5] != b[5]:
            continue
        ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
        oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
        if ox > 0.04 and oy > 0.04:
            hard.append(f"[S{a[5]}] 文本框重叠 {ox:.2f}×{oy:.2f} | {a[4]!r} <-> {b[4]!r}")

    best, run = 0, 1
    for si in range(2, len(prs.slides) + 1):
        run = run + 1 if header_fp.get(si) and header_fp[si] == header_fp.get(si - 1) else 1
        best = max(best, run)
    if best >= 4:
        warn.append(f"[deck] 连续 {best} 页共用同一页眉骨架 —— 版式雷同，AI 感首要来源")
    nb = sum(1 for v in band_pages.values() if v)
    if nb >= 3:
        warn.append(f"[deck] {nb} 页有全宽页眉/页脚色带 —— 除非它是所声明来源的结构组件"
                    "且承载内容，否则是模板脸（pitfalls #17）")
    bare_run, bare_best = 0, 0
    for si in range(1, len(prs.slides) + 1):
        bare_run = bare_run + 1 if bare_pages.get(si) else 0
        bare_best = max(bare_best, bare_run)
    if bare_best >= 3 and len(prs.slides) >= 3:
        warn.append(f"[deck] 连续 {bare_best} 页既无图片也无图表 —— 视觉载体断档，"
                    "纯文字+色块堆叠是 AI 脸结构信号（styles.md §6.5）")

    # Pillow 缺席或非 Windows 无字体目录 = 整类度量降级，走 degraded（报告尾部组装）；
    # 只有「目录在、个别字体缺映射」才是普通告警。
    if _missing and _FONT_DIR.is_dir():
        warn.append(f"[字体] {sorted(_missing)} 没有度量映射或本机缺失 —— "
                    "宽度回落估算，结果偏保守")
    unsafe = sorted(f for f in faces
                    if f and f != "?" and f.strip().lower() not in _SAFE_FONTS)
    if unsafe:
        warn.append(f"[字体] {unsafe} 不在跨平台安全名单内 —— 发给 Mac 等机器会被替换")

    for lvl, msg in check_delivery(name):
        (hard if lvl == "硬伤" else warn).append(f"[交付] {msg}")

    if WANT_PALETTE:
        for si, slide in enumerate(prs.slides, 1):
            areas: dict[str, float] = {}
            try:
                bg = _solid_rgb(slide.background.fill)
                if bg is not None:
                    areas[f"#{bg:06X}"] = SW * SH
            except Exception:
                pass
            carriers = 0
            for sh in slide.shapes:
                st = getattr(sh, "shape_type", None)
                if st == 13 or getattr(sh, "has_chart", False):
                    carriers += 1
                    continue
                rgb = _solid_rgb(_fill_of(sh))
                if rgb is not None:
                    a = (sh.width / EMU) * (sh.height / EMU)
                    areas[f"#{rgb:06X}"] = areas.get(f"#{rgb:06X}", 0.0) + a
            total = SW * SH
            top = sorted(areas.items(), key=lambda kv: -kv[1])[:4]
            rep = " · ".join(f"{c} {a / total * 100:.0f}%" for c, a in top)
            extra = f" · 图/表×{carriers}" if carriers else ""
            print(f"[palette] S{si}: {rep}{extra}  （近似值，重叠形状会重复计入）")

    if _pil_missing:
        metrics = "PIL 缺失，全部估算"
    elif _missing:
        metrics = "部分回落估算"
    else:
        metrics = "真实字体"
    degraded = list(_DEGRADED)
    if _pil_missing:
        degraded.append("字体真实度量：Pillow 未安装（python -m pip install pillow）"
                        "—— 宽度全部回落估算")
    elif _missing and not _FONT_DIR.is_dir():
        degraded.append("字体真实度量：本机没有 Windows 字体目录（非 Windows）"
                        "—— 宽度全部回落估算")
    print("\n".join(_report_lines(hard, warn, degraded, len(texts), metrics)))
    return 1 if hard else (2 if (warn or degraded) else 0)

if __name__ == "__main__":
    sys.exit(main())
