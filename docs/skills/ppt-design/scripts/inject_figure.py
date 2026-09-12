#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""把中文声明式「插图规格」编译成高分辨率透明图，嵌进 pptx 的锚点区域。

给考试卷/教材类主题补两类图：
  函数图像 —— 坐标轴、曲线、关键点、渐近线、积分阴影（matplotlib 主场，零新依赖）；
  平面几何 —— 三角形/圆/弧/角标记/直角/等长刻度/平行箭头/辅助虚线/引线标注。

设计原则：机械层收归编译器（线宽、字体、标签偏移、标记样式、颜色角色、DPI），
表达层留给 AI（画哪些点、什么关系、怎么标）。**词表 + 逃生舱两层**，绝不封死：
高频图走原语词表（可验算几何约束、风格天然统一）；词表覆盖不到的怪图，允许在规格里
写一段受限的 matplotlib 片段（"自定义代码"）兜底——仍强制统一字体/颜色/DPI/落位，
但会在报告与验证里标注「未受约束验算」。

几何正确性是本脚本的核心价值：AI 必须把自己断言的几何关系写进 "约束"，
编译器逐条验算（等长/直角/平行/共线/中点/圆过点/角度/长度），超差即报错中止、
文件保持原样。裸写绘图代码根本没有这道保险——算错一个坐标没人知道。

用法：
  python inject_figure.py deck.pptx figures.json --colors "INK=111111,ACCENT=E4002B,MUTED=6E6E6E"
  python inject_figure.py deck.pptx figures.json --dpi 600
  python inject_figure.py deck.pptx figures.json --replace        # 先清已有 figure:* 再注入
  python inject_figure.py deck.pptx figures.json --force          # 约束超差也照渲染（会大声警告）
  python inject_figure.py deck.pptx figures.json --dump-png out/  # 落一份 PNG 给人眼/judge 看
  python inject_figure.py deck.pptx figures.json --no-verify

规格格式（中文键；锚点 = 施工时的 objectName 占位矩形）：
  {
    "pages": {
      "3": [
        {"锚点": "图区", "适配": "适应", "锚点处理": "保留", "颜色": "INK",
         "字体": {"中文": "微软雅黑", "数学": "cm"},
         "图": {
           "坐标": {"x": [-1, 5], "y": [-6, 4]}, "等比例": true,
           "坐标轴": {"显示": true, "箭头": true, "刻度": true, "网格": false, "原点": "O"},
           "点": [{"名": "A", "坐标": [0, 0]}],
           "线段": [{"点": ["A", "B"], "样式": "虚线", "箭头": true}],
           "多边形": [{"顶点": ["A", "B", "C"], "填充": true, "透明": 0.12}],
           "圆": [{"心": "O", "半径": 2}], "弧": [{"心": "O", "半径": 2, "起角": 0, "终角": 90}],
           "角标记": [{"顶点": "A", "边": ["B", "C"], "弧数": 1}],
           "直角": [{"顶点": "B", "边": ["A", "C"]}],
           "等长": [{"线段": [["A", "B"], ["A", "C"]], "刻度数": 1}],
           "平行": [{"线段": [["A", "B"], ["C", "D"]]}],
           "函数": [{"表达式": "x^2-2*x-3", "定义域": [-1, 4], "颜色": "ACCENT"}],
           "阴影": [{"函数": 0, "区间": [0, 2], "颜色": "ACCENT", "透明": 0.15}],
           "标注": [{"文字": "∠A=36°", "引用": "A", "偏移": [0.3, -0.4], "引线": true}],
           "约束": [{"类型": "直角", "顶点": "A", "边": ["B", "C"]},
                    {"类型": "等长", "对象": [["A", "B"], ["A", "C"]]}],
           "标题": "图 1", "图注": "单位：cm"
         }}]}}
  适配：适应（等比内含，默认）/ 充满宽度 / 拉伸（会变形，等比例图不行）
  锚点处理：保留（默认，占位本身是面板）/ 删除（纯占位框）
  颜色：--colors 里的角色名（缺省内置 INK/ACCENT/MUTED），或六位十六进制
  数学片段用 $...$（走 cm 数学字形），汉字交给 字体.中文——别把 ^ 写在 $ 外面
  表达式里 ^ 会自动当幂运算；可用 sin/cos/tan/exp/log/sqrt/abs/pi/e 等（log 是自然对数）

验证：三层——重开计数（figure:*）+ 逐图有效 PPI + 约束报告；COM 兼容打开并打印应用身份
（12.x = WPS 兼容层，验证强度弱于真 PowerPoint）；--dump-png 落图供人眼/judge。
防重复：插图统一命名 figure:标签；重复注入会被拦，需 --replace。
退出码：0 成功；1 验证不一致；2 跳过验证；规格/约束错误直接报错、文件不动。
"""
import json
import math
import os
import re
import shutil
import sys
import tempfile

FITS = ("适应", "充满宽度", "拉伸")
STYLES = {"实线": "-", "虚线": "--", "点线": ":", "点划线": "-."}
DEFAULT_ROLES = {"INK": "111111", "ACCENT": "E4002B", "MUTED": "6E6E6E"}
TOL_LEN = 0.01          # 长度/等长/中点/圆过点：相对差
TOL_ORTHO = 0.02        # 直角 |cos| / 平行 |sin| / 共线归一化面积
TOL_ANGLE = 1.0         # 角度：绝对差（度）
_SAFE_NAMES = ("sin", "cos", "tan", "asin", "acos", "atan", "arctan", "sinh", "cosh",
               "tanh", "exp", "log", "log2", "log10", "sqrt", "cbrt", "abs", "floor",
               "ceil", "sign", "power", "minimum", "maximum", "where", "pi", "e")
_BANNED = re.compile(r"__|import|open\(|eval\(|exec\(|globals|locals|getattr|setattr|"
                     r"compile|input|os\.|sys\.|subprocess|lambda|;|`")
_EXEC_BUILTINS = {"range": range, "len": len, "min": min, "max": max, "abs": abs,
                  "enumerate": enumerate, "zip": zip, "list": list, "dict": dict,
                  "sum": sum, "float": float, "int": int, "sorted": sorted,
                  "round": round, "print": print, "tuple": tuple}


class FigureError(Exception):
    """规格/渲染/约束错误——报错中止，文件保持原样。"""


# ------------------------------------------------------------------ 向量小工具

def _sub(a, b):
    return (a[0] - b[0], a[1] - b[1])

def _add(a, b):
    return (a[0] + b[0], a[1] + b[1])

def _mul(a, k):
    return (a[0] * k, a[1] * k)

def _dot(a, b):
    return a[0] * b[0] + a[1] * b[1]

def _cross(a, b):
    return a[0] * b[1] - a[1] * b[0]

def _norm(a):
    return math.hypot(a[0], a[1])

def _unit(a):
    n = _norm(a)
    return (a[0] / n, a[1] / n) if n > 1e-12 else (0.0, 0.0)

def _ref(ref, pts):
    """点名或坐标 -> (x, y)。点名不存在直接抛错（绝不静默画错）。"""
    if isinstance(ref, str):
        if ref not in pts:
            raise FigureError(f"引用了不存在的点「{ref}」；已定义：{'、'.join(pts) or '（无）'}")
        return pts[ref]
    return (float(ref[0]), float(ref[1]))

def _seg(ref, pts):
    if isinstance(ref, str) and ref not in pts and "-" in ref:   # "A-B" 写法
        a, b = ref.split("-", 1)
        return _ref(a, pts), _ref(b, pts)
    if isinstance(ref, (list, tuple)) and len(ref) == 2 and not isinstance(ref[0], (int, float)):
        return _ref(ref[0], pts), _ref(ref[1], pts)
    raise FigureError(f"线段写法不对：{ref!r}（用 [\"A\",\"B\"] 或 \"A-B\"）")

def _pt_seg_dist(p, a, b):
    ab = _sub(b, a)
    denom = _dot(ab, ab)
    if denom < 1e-12:
        return _norm(_sub(p, a))
    t = max(0.0, min(1.0, _dot(_sub(p, a), ab) / denom))
    return _norm(_sub(p, _add(a, _mul(ab, t))))

def _nice_step(rng):
    if rng <= 0:
        return 1.0
    raw = rng / 8.0
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return m * mag
    return 10 * mag


# ------------------------------------------------------------------ 颜色 / 字体

def parse_colors(arg):
    out = {}
    for part in (arg or "").split(","):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip().lstrip("#").upper()
    return out

def _color_fn(roles, default_role):
    def c(raw=None):
        name = str(raw if raw not in (None, "") else default_role).lstrip("#").upper()
        if name in roles:
            return "#" + roles[name]
        if len(name) == 6 and all(ch in "0123456789ABCDEF" for ch in name):
            return "#" + name
        raise FigureError(f"颜色「{raw}」既不是角色（{'、'.join(roles)}）也不是六位十六进制")
    return c

_FONT_ALIAS = {"微软雅黑": "Microsoft YaHei", "黑体": "SimHei", "宋体": "SimSun",
               "等线": "DengXian", "楷体": "KaiTi", "仿宋": "FangSong",
               "思源黑体": "Source Han Sans SC", "苹方": "PingFang SC"}

def _cjk_font(name, warn):
    want = _FONT_ALIAS.get(name, name)
    try:
        from matplotlib import font_manager as fm
    except Exception:
        return want
    avail = {f.name for f in fm.fontManager.ttflist}
    if want in avail:
        return want
    for cand in ("Microsoft YaHei", "SimHei", "SimSun", "Noto Sans CJK SC",
                 "Source Han Sans SC", "PingFang SC", "DengXian", "Microsoft JhengHei"):
        if cand in avail:
            warn.append(f"字体「{name}」本机没有，中文回退到「{cand}」")
            return cand
    warn.append(f"没找到中文字体（{name}/微软雅黑/黑体/宋体都没匹配上）——"
                "中文可能渲染成方块，装字体或改 字体.中文")
    return want


# ------------------------------------------------------------------ 表达式安全求值

def _make_func(expr):
    src = str(expr).replace("^", "**")
    if _BANNED.search(src):
        raise FigureError(f"函数表达式里有不允许的内容：{expr!r}")
    if len(src) > 300:
        raise FigureError("函数表达式太长（>300 字符）")
    code = compile(src, "<figure-expr>", "eval")

    def f(xs):
        ns = {"x": xs, "np": np}
        for n in _SAFE_NAMES:
            ns[n] = getattr(np, n)
        with np.errstate(all="ignore"):
            return eval(code, {"__builtins__": {}}, ns)   # noqa: S307 - 白名单命名空间
    return f

def _sample(expr, xs, warn=None):
    """求值并取实部——log/sqrt 的负数域会变 NaN（曲线留空档），复数只取实部。"""
    raw = np.asarray(_make_func(expr)(xs))
    if np.iscomplexobj(raw):
        if warn is not None:
            warn.append(f"函数 {expr!r} 有复数结果，只画了实部（定义域检查一下）")
        raw = np.real(raw)
    return np.asarray(raw, dtype=float)


# ------------------------------------------------------------------ 约束验算

def check_constraints(constraints, pts, circles):
    """返回 (fails, warns)；每条是一行中文报告。等长/直角/平行/共线/中点/圆过点/角度/长度。"""
    fails, warns = [], []
    for i, con in enumerate(constraints or []):
        if not isinstance(con, dict):
            fails.append(f"约束{i + 1} 不是对象：{con!r}")
            continue
        t = con.get("类型")
        try:
            if t == "等长":
                segs = [_seg(s, pts) for s in (con.get("对象") or con.get("线段"))]
                lens = [_norm(_sub(b, a)) for a, b in segs]
                if min(lens) <= 1e-9:
                    fails.append(f"约束 等长：存在零长线段（{lens}）")
                else:
                    dev = (max(lens) - min(lens)) / max(lens)
                    line = (f"约束 等长（{len(segs)} 段）：实测 {' / '.join(f'{v:.3f}' for v in lens)}"
                            f"（最大差 {dev:.1%}）")
                    (warns if dev <= TOL_LEN else fails).append(line + ("" if dev <= TOL_LEN else " → 超差"))
            elif t == "直角":
                v = _ref(con["顶点"], pts)
                b, d = (_ref(x, pts) for x in con["边"])
                cosv = abs(_dot(_unit(_sub(b, v)), _unit(_sub(d, v))))
                ok = cosv <= TOL_ORTHO
                line = (f"约束 直角（顶点 {con['顶点']}）：|cos|={cosv:.4f}"
                        f"（≈{math.degrees(math.acos(min(1.0, cosv))):.2f}°）")
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            elif t == "平行":
                s1, s2 = (_seg(s, pts) for s in con["线段"])
                sinv = abs(_cross(_unit(_sub(s1[1], s1[0])), _unit(_sub(s2[1], s2[0]))))
                ok = sinv <= TOL_ORTHO
                line = f"约束 平行：|sin|={sinv:.4f}"
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            elif t == "共线":
                a, b, d = (_ref(x, pts) for x in con["点"])
                u, w = _sub(b, a), _sub(d, a)
                denom = _norm(u) * _norm(w)
                dev = abs(_cross(u, w)) / denom if denom > 1e-12 else 1.0
                ok = dev <= TOL_ORTHO
                line = f"约束 共线（{con['点']}）：偏离 {dev:.4f}"
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            elif t == "中点":
                m = _ref(con["中点"], pts)
                a, b = _ref(con["线段"][0], pts), _ref(con["线段"][1], pts)
                err = _norm(_sub(m, _mul(_add(a, b), 0.5))) / max(_norm(_sub(b, a)), 1e-9)
                ok = err <= TOL_LEN
                line = (f"约束 中点（{con['中点']} = {con['线段'][0]}{con['线段'][1]} 中点）："
                        f"偏差 {err:.2%}")
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            elif t == "圆过点":
                if "圆" in con:
                    cen, r = circles[int(con["圆"])]["心"], circles[int(con["圆"])]["半径"]
                else:
                    cen, r = _ref(con["心"], pts), float(con["半径"])
                for pn in con["点"]:
                    d = _norm(_sub(_ref(pn, pts), cen))
                    err = abs(d - r) / max(r, 1e-9)
                    line = (f"约束 圆过点（点 {pn} 到心 {d:.3f}，半径 {r:.3f}）：偏差 {err:.2%}")
                    (warns if err <= TOL_LEN else fails).append(line + ("" if err <= TOL_LEN else " → 超差"))
            elif t == "角度":
                v = _ref(con["顶点"], pts)
                b, d = (_ref(x, pts) for x in con["边"])
                got = math.degrees(math.acos(max(-1.0, min(1.0,
                          _dot(_unit(_sub(b, v)), _unit(_sub(d, v)))))))
                want = float(con["角度"])
                ok = abs(got - want) <= TOL_ANGLE
                line = f"约束 角度（顶点 {con['顶点']}）：声明 {want:g}°，实测 {got:.2f}°"
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            elif t == "长度":
                a, b = _seg(con["线段"], pts)
                got, want = _norm(_sub(b, a)), float(con["长度"])
                err = abs(got - want) / max(want, 1e-9)
                ok = err <= TOL_LEN
                line = f"约束 长度（{con['线段']}）：声明 {want:g}，实测 {got:.3f}（差 {err:.2%}）"
                (warns if ok else fails).append(line + ("" if ok else " → 超差"))
            else:
                fails.append(f"约束类型「{t}」不认识；可用：等长/直角/平行/共线/中点/圆过点/角度/长度")
        except FigureError as e:
            fails.append(f"约束{i + 1}（{t}）：{e}")
        except Exception as e:
            fails.append(f"约束{i + 1}（{t}）无法验算：{str(e)[:90]}")
    return fails, warns


# ------------------------------------------------------------------ 绘图

def _ticks(lo, hi, step):
    out, t, k = [], math.ceil(lo / step) * step, 0
    while t <= hi + 1e-9 and k < 200:
        out.append(0.0 if abs(t) < 1e-9 else round(t, 10))
        t += step
        k += 1
    return out

def _fmt(v):
    return f"{v:g}"

def _draw_axes(ax, axd, xr, yr, c):
    ink = c()
    if not axd.get("显示", True):
        return
    x0, x1 = xr
    y0, y1 = yr
    ybase = 0.0 if y0 <= 0 <= y1 else (y0 if y0 > 0 else y1)
    xbase = 0.0 if x0 <= 0 <= x1 else (x0 if x0 > 0 else x1)
    span = max(x1 - x0, y1 - y0)
    arrow = axd.get("箭头", True)
    if arrow:
        ax.annotate("", xy=(x1, ybase), xytext=(x0, ybase), zorder=2,
                    arrowprops=dict(arrowstyle="-|>", color=ink, lw=1.1,
                                    shrinkA=0, shrinkB=0, mutation_scale=12))
        ax.annotate("", xy=(xbase, y1), xytext=(xbase, y0), zorder=2,
                    arrowprops=dict(arrowstyle="-|>", color=ink, lw=1.1,
                                    shrinkA=0, shrinkB=0, mutation_scale=12))
    else:
        ax.plot([x0, x1], [ybase, ybase], color=ink, lw=1.1, zorder=2)
        ax.plot([xbase, xbase], [y0, y1], color=ink, lw=1.1, zorder=2)
    step = _nice_step(max(x1 - x0, y1 - y0))
    tick = 0.014 * span
    if axd.get("网格", False):
        for tx in _ticks(x0, x1, step):
            ax.plot([tx, tx], [y0, y1], color=c("MUTED"), lw=0.6, ls=":", alpha=0.5, zorder=0)
        for ty in _ticks(y0, y1, step):
            ax.plot([x0, x1], [ty, ty], color=c("MUTED"), lw=0.6, ls=":", alpha=0.5, zorder=0)
    if axd.get("刻度", True):
        tl = axd.get("刻度字号", 9)
        for tx in _ticks(x0, x1, step):
            if abs(tx) < 1e-9 and axd.get("原点"):
                continue
            ax.plot([tx, tx], [ybase - tick, ybase + tick], color=ink, lw=0.9, zorder=2)
            ax.text(tx, ybase - tick * 2.1, _fmt(tx), ha="center", va="top",
                    fontsize=tl, color=ink, zorder=6)
        for ty in _ticks(y0, y1, step):
            if abs(ty) < 1e-9 and axd.get("原点"):
                continue
            ax.plot([xbase - tick, xbase + tick], [ty, ty], color=ink, lw=0.9, zorder=2)
            ax.text(xbase - tick * 2.1, ty, _fmt(ty), ha="right", va="center",
                    fontsize=tl, color=ink, zorder=6)
    if axd.get("原点"):
        ax.text(xbase - tick * 2.1, ybase - tick * 2.1, str(axd["原点"]),
                ha="right", va="top", fontsize=axd.get("刻度字号", 9) + 1, color=ink, zorder=6)
    if axd.get("x标签"):
        ax.text(x1, ybase - tick * 2.1, " " + str(axd["x标签"]), ha="left", va="top",
                fontsize=11, color=ink, zorder=6)
    if axd.get("y标签"):
        ax.text(xbase + tick * 1.6, y1, " " + str(axd["y标签"]), ha="left", va="center",
                fontsize=11, color=ink, zorder=6)


def _angle_sweep(v, b, d):
    a1 = math.degrees(math.atan2(b[1] - v[1], b[0] - v[0])) % 360
    a2 = math.degrees(math.atan2(d[1] - v[1], d[0] - v[0])) % 360
    if (a2 - a1) % 360 > 180:
        a1, a2 = a2, a1
    return a1, a2


def _range_of(fd, pts):
    xs, ys = [], []
    for p in pts.values():
        xs.append(p[0]); ys.append(p[1])
    for s in fd.get("线段", []) or []:
        a, b = _seg(s["点"], pts)
        xs += [a[0], b[0]]; ys += [a[1], b[1]]
    for poly in fd.get("多边形", []) or []:
        for v in poly["顶点"]:
            x, y = _ref(v, pts); xs.append(x); ys.append(y)
    for key in ("圆", "弧"):
        for item in fd.get(key, []) or []:
            cen = _ref(item["心"], pts)
            r = (float(item["半径"]) if "半径" in item
                 else _norm(_sub(_ref(item["过点"][0], pts), cen)))
            xs += [cen[0] - r, cen[0] + r]; ys += [cen[1] - r, cen[1] + r]
    rng = fd.get("坐标")
    xr = [float(v) for v in rng["x"]] if rng and "x" in rng else None
    yr = [float(v) for v in rng["y"]] if rng and "y" in rng else None

    def span(vals, given):
        if given:
            return given
        if not vals:
            return [-1.0, 1.0]
        lo, hi = min(vals), max(vals)
        pad = (hi - lo) * 0.08 or 1.0
        return [lo - pad, hi + pad]

    xr, yr = span(xs, xr), span(ys, yr)
    if xr[1] - xr[0] <= 1e-9 or yr[1] - yr[0] <= 1e-9:
        raise FigureError("坐标范围退化（宽或高为 0）——检查点坐标或 坐标 设置")
    return (xr[0], xr[1]), (yr[0], yr[1])


def _draw(ax, fd, c, warn):
    """把「图」字典画到 ax 上；返回 [(xs, ys), ...] 曲线，供阴影引用。"""
    from matplotlib import pyplot as plt
    from matplotlib.patches import Arc, Circle, Polygon as MplPoly
    pts, curves = {}, []
    for p in fd.get("点", []) or []:
        if not p.get("名"):
            raise FigureError("点必须给 名（线段/标记要用它引用）")
        if p["名"] in pts:
            raise FigureError(f"点「{p['名']}」重复定义")
        pts[p["名"]] = (float(p["坐标"][0]), float(p["坐标"][1]))

    xr, yr = _range_of(fd, pts)
    geo = any(fd.get(k) for k in ("圆", "弧", "角标记", "直角", "等长", "多边形"))
    equal = fd.get("等比例")
    if equal is None:
        equal = bool(geo and not fd.get("函数"))
    equal = bool(equal)
    if geo and fd.get("等比例") is False:
        warn.append("图里有几何图元但 等比例 显式设了 false——圆/角/直角会看起来变形")

    if equal:
        ax.set_aspect("equal", adjustable="box")
    ax.set_xlim(*xr)
    ax.set_ylim(*yr)
    ax.axis("off")
    axd = fd.get("坐标轴", {})
    if axd is True:
        axd = {"显示": True}
    if axd:
        _draw_axes(ax, axd, xr, yr, c)
    span = max(xr[1] - xr[0], yr[1] - yr[0])

    # 函数曲线
    for f in fd.get("函数", []) or []:
        n = int(f.get("采样", 400))
        dom = f.get("定义域") or [xr[0], xr[1]]
        xs = np.linspace(float(dom[0]), float(dom[1]), n)
        ys = _sample(f["表达式"], xs, warn)
        bad = int(np.sum(~np.isfinite(ys)))
        if bad > n * 0.3:
            warn.append(f"函数 {f['表达式']!r} 在定义域内有 {bad}/{n} 个点无实数值"
                        "（log/sqrt 负数？）——曲线会断成空档")
        ax.plot(xs, ys, color=c(f.get("颜色")), lw=float(f.get("宽", 1.6)),
                ls=STYLES.get(f.get("样式", "实线"), "-"), zorder=3)
        curves.append((xs, ys))

    # 阴影（zorder 压在曲线之下）
    for sh in fd.get("阴影", []) or []:
        alpha, color = float(sh.get("透明", 0.15)), c(sh.get("颜色", "ACCENT"))
        if "多边形" in sh:
            ax.add_patch(MplPoly([_ref(v, pts) for v in sh["多边形"]], closed=True,
                                 facecolor=color, edgecolor="none", alpha=alpha, zorder=1.5))
            continue
        a, b = (float(v) for v in sh["区间"])
        xs = np.linspace(a, b, 400)
        idx = sh.get("函数")
        if isinstance(idx, (list, tuple)) and len(idx) == 2:
            y1 = _sample(fd["函数"][int(idx[0])]["表达式"], xs, warn)
            y2 = _sample(fd["函数"][int(idx[1])]["表达式"], xs, warn)
            ax.fill_between(xs, y1, y2, color=color, alpha=alpha, lw=0, zorder=1.5)
        else:
            ys = _sample(fd["函数"][int(idx)]["表达式"], xs, warn)
            ax.fill_between(xs, ys, float(sh.get("基准", 0.0)), color=color, alpha=alpha,
                            lw=0, zorder=1.5)

    for poly in fd.get("多边形", []) or []:
        filled = bool(poly.get("填充", False))
        ax.add_patch(MplPoly([_ref(v, pts) for v in poly["顶点"]], closed=True,
                             facecolor=c(poly.get("填充色", poly.get("颜色"))) if filled else "none",
                             edgecolor=c(poly.get("颜色")), lw=float(poly.get("宽", 1.5)),
                             alpha=float(poly.get("透明", 1.0)) if filled else 1.0, zorder=2.5))

    circles = []
    for ci in fd.get("圆", []) or []:
        cen = _ref(ci["心"], pts)
        r = (float(ci["半径"]) if "半径" in ci
             else _norm(_sub(_ref(ci["过点"][0], pts), cen)))
        circles.append({"心": cen, "半径": r})
        ax.add_patch(Circle(cen, r, facecolor="none", edgecolor=c(ci.get("颜色")),
                            lw=float(ci.get("宽", 1.5)),
                            ls=STYLES.get(ci.get("样式", "实线"), "-"), zorder=2))

    for a in fd.get("弧", []) or []:
        cen = _ref(a["心"], pts)
        r = (float(a["半径"]) if "半径" in a
             else _norm(_sub(_ref(a["过点"][0], pts), cen)))
        ax.add_patch(Arc(cen, 2 * r, 2 * r, angle=0, theta1=float(a.get("起角", 0)),
                         theta2=float(a.get("终角", 90)), edgecolor=c(a.get("颜色")),
                         lw=float(a.get("宽", 1.5)),
                         ls=STYLES.get(a.get("样式", "实线"), "-"), zorder=2))

    for ang in fd.get("角标记", []) or []:
        v = _ref(ang["顶点"], pts)
        b, d = (_ref(x, pts) for x in ang["边"])
        a1, a2 = _angle_sweep(v, b, d)
        r = float(ang["半径"]) if ang.get("半径") else span * 0.07
        for k in range(int(ang.get("弧数", 1))):
            rr = r + k * span * 0.022
            ax.add_patch(Arc(v, 2 * rr, 2 * rr, angle=0, theta1=a1, theta2=a2,
                             edgecolor=c(ang.get("颜色")), lw=1.3, zorder=4))

    for rt in fd.get("直角", []) or []:
        v = _ref(rt["顶点"], pts)
        b, d = (_ref(x, pts) for x in rt["边"])
        s = float(rt["边长"]) if rt.get("边长") else span * 0.045
        p1, p3 = _add(v, _mul(_unit(_sub(b, v)), s)), _add(v, _mul(_unit(_sub(d, v)), s))
        p2 = _add(p1, _mul(_unit(_sub(d, v)), s))
        ax.plot([p1[0], p2[0], p3[0]], [p1[1], p2[1], p3[1]],
                color=c(rt.get("颜色")), lw=1.3, zorder=4)

    for eq in fd.get("等长", []) or []:
        n = int(eq.get("刻度数", 1))
        for a, b in [_seg(s, pts) for s in eq["线段"]]:
            mid, along = _mul(_add(a, b), 0.5), _unit(_sub(b, a))
            perp, hl = (-along[1], along[0]), span * 0.018
            for k in range(n):
                ctr = _add(mid, _mul(along, (k - (n - 1) / 2) * hl * 1.4))
                p, q = _add(ctr, _mul(perp, hl)), _add(ctr, _mul(perp, -hl))
                ax.plot([p[0], q[0]], [p[1], q[1]], color=c(eq.get("颜色")), lw=1.3, zorder=4)

    for pa in fd.get("平行", []) or []:
        for a, b in [_seg(s, pts) for s in pa["线段"]]:
            mid, along = _mul(_add(a, b), 0.5), _unit(_sub(b, a))
            perp, s = (-along[1], along[0]), span * 0.022
            tip = _add(mid, _mul(along, s))
            t1 = _add(_sub(tip, _mul(along, s)), _mul(perp, s * 0.7))
            t2 = _add(_sub(tip, _mul(along, s)), _mul(perp, -s * 0.7))
            ax.plot([t1[0], tip[0], t2[0]], [t1[1], tip[1], t2[1]],
                    color=c(pa.get("颜色")), lw=1.3, zorder=4)

    for s in fd.get("线段", []) or []:
        a, b = _seg(s["点"], pts)
        col, lw, ls = c(s.get("颜色")), float(s.get("宽", 1.5)), STYLES.get(s.get("样式", "实线"), "-")
        if s.get("箭头", False):
            ax.annotate("", xy=b, xytext=a, zorder=3,
                        arrowprops=dict(arrowstyle="<|-|>" if s["箭头"] == "双向" else "-|>",
                                        color=col, lw=lw, linestyle=ls, shrinkA=0, shrinkB=0,
                                        mutation_scale=13))
        else:
            ax.plot([a[0], b[0]], [a[1], b[1]], color=col, lw=lw, ls=ls, zorder=3)

    # 点与点标签
    allpts = list(pts.values())
    cen = ((sum(p[0] for p in allpts) / len(allpts), sum(p[1] for p in allpts) / len(allpts))
           if allpts else (0.0, 0.0))
    segs_geom = [_seg(s["点"], pts) for s in fd.get("线段", []) or []]
    for poly in fd.get("多边形", []) or []:
        vv = [_ref(v, pts) for v in poly["顶点"]]
        segs_geom += [(vv[k], vv[(k + 1) % len(vv)]) for k in range(len(vv))]
    for p in fd.get("点", []) or []:
        xy, col = pts[p["名"]], c(p.get("颜色"))
        if p.get("标记", "圆点") != "无":
            ax.plot([xy[0]], [xy[1]], marker="o", ms=float(p.get("点大小", 4.2)),
                    color=col, zorder=5)
        if not p.get("显示标注", True):
            continue
        if p.get("标注偏移"):
            d = (float(p["标注偏移"][0]), float(p["标注偏移"][1]))
        else:
            d = _unit(_sub(xy, cen))
            if d == (0.0, 0.0):
                d = _unit((0.7, 0.7))
            d = _mul(d, span * 0.045)
        tx, ty = xy[0] + d[0], xy[1] + d[1]
        for a, b in segs_geom:
            if p["名"] in (a, b) or _norm(_sub(a, xy)) < 1e-9 or _norm(_sub(b, xy)) < 1e-9:
                continue
            if _pt_seg_dist((tx, ty), a, b) < span * 0.025:
                warn.append(f"点「{p['名']}」的标签可能压在线上——给「标注偏移」挪开")
                break
        ax.text(tx, ty, str(p["名"]), fontsize=p.get("字号", 12), color=col, zorder=6,
                ha="left" if d[0] >= 0 else "right", va="bottom" if d[1] >= 0 else "top")

    for an in fd.get("标注", []) or []:
        pos = _ref(an["引用"], pts) if "引用" in an else _ref(an.get("坐标", [0, 0]), pts)
        off = an.get("偏移", [span * 0.05, span * 0.05])
        tgt = (pos[0] + float(off[0]), pos[1] + float(off[1]))
        # 只看锚点是否落在线上（低噪提示）；文字横穿线条的观感交闸门二 judge
        for a, b in segs_geom:
            if _pt_seg_dist(tgt, a, b) < span * 0.03:
                warn.append(f"标注「{str(an['文字'])[:14]}」离线条很近——核对是否压线，必要时挪「偏移」")
                break
        if an.get("引线"):
            ax.annotate(str(an["文字"]), xy=pos, xytext=tgt, zorder=6,
                        fontsize=an.get("字号", 11), color=c(an.get("颜色")),
                        ha=an.get("对齐", "left"), va="bottom",
                        arrowprops=dict(arrowstyle="-", color=c(an.get("颜色")), lw=0.9))
        else:
            ax.text(tgt[0], tgt[1], str(an["文字"]), fontsize=an.get("字号", 11),
                    color=c(an.get("颜色")), zorder=6, ha=an.get("对齐", "left"), va="bottom")

    if fd.get("标题"):
        ax.set_title(str(fd["标题"]), fontsize=13, color=c(fd.get("标题颜色", "INK")), pad=6)
    if fd.get("图注"):
        ax.text(0.5, -0.09, str(fd["图注"]), transform=ax.transAxes, ha="center", va="top",
                fontsize=9, color=c("MUTED"))

    if fd.get("自定义代码"):
        warn.append("本图含「自定义代码」——已按 AI 代码绘制，未受几何约束验算")
        ns = {"ax": ax, "np": np, "plt": plt, "c": c, "__builtins__": _EXEC_BUILTINS}
        try:
            exec(compile(str(fd["自定义代码"]), "<figure-custom>", "exec"), ns)  # noqa: S102
        except Exception as e:
            raise FigureError(f"自定义代码执行失败：{str(e)[:120]}")
    return curves


# ------------------------------------------------------------------ 渲染（按目标英寸）

def _build(figsize, fd, c, warn, fonts):
    from matplotlib import pyplot as plt
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": [fonts["中文"], "SimHei", "Microsoft YaHei", "DejaVu Sans"],
        "axes.unicode_minus": False,          # 防负号变方块
        "mathtext.fontset": fonts["数学"],
        "figure.dpi": 100,
    })
    fig = plt.figure(figsize=figsize)
    ax = fig.add_axes([0.035, 0.035, 0.93, 0.93])
    _draw(ax, fd, c, warn)
    return fig


def _tight(fig, dpi, path, bg):
    # 注意：matplotlib 里 transparent=True 一旦再显式传 facecolor（哪怕是 "auto"）
    # 就会失效、整张变成不透明白底（3.11 实测）。要透明就只传 transparent。
    kw = dict(dpi=dpi, bbox_inches="tight", pad_inches=0.02)
    if bg:
        kw.update(facecolor="#" + bg, transparent=False)
    else:
        kw.update(transparent=True)
    fig.savefig(path, **kw)
    from PIL import Image
    with Image.open(path) as im:
        return im.size


def render_figure(fd, c, fonts, dpi, max_w, max_h, path, bg, warn, mode="适应"):
    """量内容真实比例 → 适配进 max_w×max_h 盒子 → 按目标英寸渲染。
    返回 (px_w, px_h, in_w, in_h)。字体是绝对磅值，所以迭代 2 轮收敛到目标尺寸。"""
    from matplotlib import pyplot as plt
    probe = path + ".probe.png"

    def measure(w, h):
        fig = _build((max(w, 0.3), max(h, 0.3)), fd, c, warn, fonts)
        pw, ph = _tight(fig, 100, probe, bg)
        plt.close(fig)
        return pw / 100, ph / 100

    nat_w, nat_h = measure(4.0, 3.0)          # 中性尺寸量内容真实比例
    aspect = nat_w / nat_h if nat_h else 1.0
    if mode == "拉伸":
        want_w, want_h = max_w, max_h
    elif mode == "充满宽度":
        want_w, want_h = max_w, max_w / aspect
    elif max_w / max_h >= aspect:              # 适应：以紧的那一边为准
        want_h = max_h; want_w = max_h * aspect
    else:
        want_w = max_w; want_h = max_w / aspect
    want_w, want_h = max(want_w, 0.3), max(want_h, 0.3)
    # figsize 与「紧裁后的成品英寸」不是一回事（轴只占图的 93%），迭代让成品命中目标尺寸
    fw, fh = want_w, want_h
    for _ in range(5):
        got_w, _got_h = measure(fw, fh)
        if abs(got_w - want_w) <= 0.015 * want_w:
            break
        k = want_w / got_w if got_w else 1.0
        fw, fh = fw * k, fh * k
    fig = _build((fw, fh), fd, c, warn, fonts)
    pw, ph = _tight(fig, dpi, path, bg)
    plt.close(fig)
    try:
        os.remove(probe)
    except OSError:
        pass
    return pw, ph, pw / dpi, ph / dpi


# ------------------------------------------------------------------ 插入

def _first_run_color(sh):
    if sh is None or not getattr(sh, "has_text_frame", False):
        return None
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.color and r.font.color.rgb:
                    return str(r.font.color.rgb)
    except Exception:
        pass
    return None


def inject(src, spec, colors_arg, dpi_cli, ppi_cli, replace, force, dump_dir):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("缺依赖：python -m pip install python-pptx")
    import numpy as _np  # noqa: F401  确认渲染依赖在位（缺了给清晰错误，不半路炸）
    prs = Presentation(src)
    errors, jobs = [], []
    for si, slide in enumerate(prs.slides, 1):
        items = spec.get("pages", {}).get(str(si), [])
        if not items:
            continue
        by_name = {sh.name: sh for sh in slide.shapes if sh.name}
        for it in items:
            fd = it.get("图")
            if not isinstance(fd, dict):
                errors.append(f"S{si}: 缺「图」对象"); continue
            anchor = it.get("锚点")
            if anchor is not None and anchor not in by_name:
                errors.append(f"S{si}: 锚点形状「{anchor}」不存在；本页有："
                              + "、".join(sorted(by_name))[:120]); continue
            if anchor is None and "位置" not in it:
                errors.append(f"S{si}: 既没给锚点也没给「位置」——不知道该放哪"); continue
            if anchor is not None:
                ash = by_name[anchor]
                if None in (ash.left, ash.top, ash.width, ash.height):
                    errors.append(f"S{si}: 锚点「{anchor}」没有几何尺寸（不是可放置的形状）")
                    continue
            mode = it.get("适配", "适应")
            if mode not in FITS:
                errors.append(f"S{si}: 适配「{mode}」不合法；可选：{'、'.join(FITS)}"); continue
            jobs.append((si, slide, by_name.get(anchor), it, fd, mode))

    stale = [(si, sh) for si, slide in enumerate(prs.slides, 1) for sh in slide.shapes
             if sh.name and sh.name.startswith("figure:")]
    if stale and not replace:
        sys.exit(f"文件里已有 {len(stale)} 张插图——改规格后重注入请加 --replace")
    for _si, sh in stale:
        sh._element.getparent().remove(sh._element)
    if errors:
        for e in errors:
            print("✗ " + e)
        sys.exit("插图规格有误（文件保持原样）——修正后重跑")
    if not jobs:
        print("规格里没有可注入的插图，文件未改动")
        return 0, 0, []

    if dump_dir:
        os.makedirs(dump_dir, exist_ok=True)
    cli_colors = parse_colors(colors_arg)
    tmpdir = tempfile.mkdtemp(prefix="figure_")
    placed, warns, fails, errors = 0, [], [], []
    try:
        for i, (si, slide, anchor, it, fd, mode) in enumerate(jobs):
            label = it.get("标签") or (anchor.name if anchor is not None else f"S{si}")
            bg = str(it.get("底", "")).lstrip("#").upper() or None
            roles = dict(DEFAULT_ROLES)
            roles.update(cli_colors)
            if "INK" not in cli_colors:
                rr = _first_run_color(anchor)
                if rr:
                    roles["INK"] = rr          # 没给 --colors 时图跟随锚点文字色
            c = _color_fn(roles, str(it.get("颜色", "INK")))
            fnt = it.get("字体") or {}
            fonts = {"中文": _cjk_font(str(fnt.get("中文", "微软雅黑")), warns),
                     "数学": str(fnt.get("数学", "cm"))}
            dpi = int(it.get("dpi", dpi_cli))
            if anchor is not None:
                box_l, box_t, box_w, box_h = anchor.left, anchor.top, anchor.width, anchor.height
            else:
                p = it["位置"]
                box_l, box_t = Emu(int(float(p["左"]) * 914400)), Emu(int(float(p["上"]) * 914400))
                box_w, box_h = Emu(int(float(p["宽"]) * 914400)), Emu(int(float(p["高"]) * 914400))
            box_win, box_hin = box_w / 914400, box_h / 914400

            # 约束验算（渲染前；超差且未 --force 就不渲染这张）
            pts = {}
            for pp in fd.get("点", []) or []:
                if pp.get("名"):
                    pts[str(pp["名"])] = (float(pp["坐标"][0]), float(pp["坐标"][1]))
            circles = []
            for ci in fd.get("圆", []) or []:
                cen = _ref(ci["心"], pts)
                r = (float(ci["半径"]) if "半径" in ci
                     else _norm(_sub(_ref(ci["过点"][0], pts), cen)))
                circles.append({"心": cen, "半径": r})
            cf, co = check_constraints(fd.get("约束"), pts, circles)
            fails += [f"S{si} {label}：{x}" for x in cf]
            if co and not cf:
                print(f"  ✓ {label}：约束 {len(co)} 条验算通过")
            if cf and not force:
                continue

            png = os.path.join(tmpdir, f"g{i}.png")
            try:
                pw, ph, fwin, fhin = render_figure(fd, c, fonts, dpi, max(box_win, 0.4),
                                                   max(box_hin, 0.4), png, bg, warns, mode)
            except FigureError:
                raise
            except Exception as e:
                errors.append(f"S{si} {label}：渲染失败 —— {str(e)[:120]}")
                continue
            if mode == "拉伸":
                warns.append(f"S{si} {label}：适配=拉伸 会改变比例，圆/角度会被拉变形")
            if mode == "充满宽度" and fhin > box_hin + 1e-6:
                warns.append(f"S{si} {label}：充满宽度后高出锚点框 {fhin - box_hin:.2f}\"")
            eff = pw / max(fwin, 1e-6)
            if ppi_cli > 0 and eff > ppi_cli * 1.6:      # 用户限了 PPI 就降采样，别塞爆文件
                from PIL import Image
                keep = max(1, int(round(fwin * ppi_cli)))
                with Image.open(png) as im:
                    im.resize((keep, max(1, int(round(ph * keep / pw)))), Image.LANCZOS).save(png)
                with Image.open(png) as im:
                    pw, ph = im.size
                eff = pw / max(fwin, 1e-6)
            if eff < 150:
                warns.append(f"S{si} {label}：有效 PPI ≈{eff:.0f} 偏低——投屏会发虚")

            img_w, img_h = Emu(int(fwin * 914400)), Emu(int(fhin * 914400))
            left = max(int(box_l + (box_w - img_w) // 2), 0)
            top = max(int(box_t + (box_h - img_h) // 2), 0)
            if it.get("锚点处理", "保留") == "删除" and anchor is not None:
                anchor._element.getparent().remove(anchor._element)
            pic = slide.shapes.add_picture(png, left, top, width=img_w, height=img_h)
            pic._element.nvPicPr.cNvPr.set("name", f"figure:{label}")
            placed += 1
            tag = "（含自定义代码·未受约束验算）" if fd.get("自定义代码") else ""
            print(f"  S{si}：{label} {mode} @ {fwin:.2f}×{fhin:.2f}\" "
                  f"有效 PPI ≈{eff:.0f} → {pw}px{tag}")
            if dump_dir:
                safe = re.sub(r"[^0-9A-Za-z\u4e00-\u9fa5_-]", "_", str(label))
                shutil.copyfile(png, os.path.join(dump_dir, f"S{si}_{safe}.png"))
        if errors:
            for e in errors:
                print("✗ " + e)
            sys.exit("插图渲染失败（文件保持原样）——修正后重跑")
        if fails and not force:
            print("✗ 几何约束验算未通过：")
            for f in fails:
                print("  ✗ " + f)
            sys.exit("约束超差（文件保持原样）——按上面的实测值改坐标后重跑；确实要照画请加 --force")
        if fails:
            for f in fails:
                print("  ⚠ 约束超差（已 --force 放行）：" + f)
        prs.save(src)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    print(f"已插入 {placed} 张插图（{dpi_cli}dpi）")
    if fails:
        warns.append("有约束超差被 --force 放行——交付时要如实说明")
    warns = list(dict.fromkeys(warns))       # 定尺寸要重绘多次，逐次告警会重复
    return placed, len(jobs), warns


# ------------------------------------------------------------------ 验证

def verify(src, placed, warns):
    from pptx import Presentation
    prs = Presentation(src)
    figs = [(si, sh) for si, slide in enumerate(prs.slides, 1) for sh in slide.shapes
            if sh.name and sh.name.startswith("figure:")]
    if len(figs) != placed:
        print(f"  ✗ 重开计数 {len(figs)} ≠ 注入 {placed}")
        return False
    for si, sh in figs:
        ppi = round(sh.image.size[0] / (sh.width / 914400))
        print(f"  ✓ S{si} {sh.name} 有效 PPI ≈ {ppi}")
    for w in dict.fromkeys(warns):
        print("  ? " + w)
    try:
        import win32com.client
        app = win32com.client.dynamic.Dispatch("PowerPoint.Application")
        try:
            app.DisplayAlerts = 1
        except Exception:
            pass
        pres = app.Presentations.Open(os.path.abspath(src), ReadOnly=True, WithWindow=False)
        pres.Close()
        try:
            if app.Presentations.Count == 0:
                app.Quit()
        except Exception:
            pass
        note = ("；12.x = WPS 兼容层，验证强度弱于真 PowerPoint"
                if str(app.Version).startswith("12.") else "")
        print(f"  ✓ PowerPoint 兼容 COM 打开无修复（{app.Name} {app.Version}.{app.Build}{note}）")
    except ImportError:
        print("  ⚠ 无 COM，跳过打开体检")
    except Exception as e:
        print(f"  ✗ COM 打开失败：{str(e)[:110]}")
        return False
    return True


# ------------------------------------------------------------------ main

def main():
    # 规范参数扫描：带值旗标吃掉它的值（否则 --colors INK=111 的值会被当成位置参数）
    # 未知旗标/给错的值一律报错退出——静默退回默认值会让 AI 以为约束已生效。
    argv = sys.argv[1:]
    known = {"--colors", "--dump-png", "--dpi", "--ppi", "--replace", "--force", "--no-verify"}
    for a in argv:
        if not a.startswith("--"):
            continue
        base, eq, _ = a.partition("=")
        if base not in known:
            sys.exit(f"未知参数 {base}；可用：{' '.join(sorted(known))}")
        if eq and base in ("--colors", "--dump-png", "--dpi", "--ppi"):
            sys.exit(f"{base} 不支持 = 形式（--colors 的值本身含 =），请用空格：{base} <值>")
    pos, colors_arg, dump_dir, dpi, ppi, flags, i = [], "", "", 600, 0, set(), 0
    while i < len(argv):
        a = argv[i]
        if a in ("--colors", "--dump-png", "--dpi", "--ppi"):
            if i + 1 >= len(argv):
                sys.exit(f"{a} 缺少值")
            val = argv[i + 1]
            if a == "--colors":
                colors_arg = val
            elif a == "--dump-png":
                dump_dir = val
            else:
                try:
                    if a == "--dpi":
                        dpi = int(val)
                    else:
                        ppi = int(val)
                except ValueError:
                    sys.exit(f"{a} 的值 {val!r} 不是整数")
            i += 2
            continue
        if a.startswith("--"):
            flags.add(a)
            i += 1
            continue
        pos.append(a)
        i += 1
    if len(pos) != 2:
        sys.exit(__doc__)
    src, spec_path = pos
    if not os.path.exists(src):
        sys.exit("找不到文件：" + src)
    try:
        spec = json.load(open(spec_path, encoding="utf-8"))
    except Exception as e:
        sys.exit(f"规格 {spec_path} 读不了：{e}")
    placed, _, warns = inject(src, spec, colors_arg, dpi, ppi,
                              replace="--replace" in flags, force="--force" in flags,
                              dump_dir=dump_dir)
    if not placed:
        return 0
    if "--no-verify" in flags:
        print("⚠ 已插入但跳过验证（--no-verify）")
        return 2
    return 0 if verify(src, placed, warns) else 1


try:
    import numpy as np
except ImportError:      # 渲染时才需要；这里给清晰错误
    np = None

if __name__ == "__main__":
    if np is None:
        sys.exit("缺依赖：python -m pip install matplotlib numpy")
    sys.exit(main())
