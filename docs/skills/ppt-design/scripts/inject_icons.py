#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""给 pptx 注入线性图标（Tabler Icons，MIT，内嵌字体，离线可用）。

图标是"路标不是装饰"：高级用法密度很低——章节标记、并列能力、流程每步各一个，
绝不是每条要点配一个。本脚本按声明式规格渲染图标为高分辨率透明图并放到锚点
形状旁边的栅格位置；单一字体家族保证笔触一致（这是字体方案相对图片素材的核心收益）。

素材与许可：assets/icons/tabler-icons.ttf + LICENSE-tabler-icons.txt（MIT，
Copyright (c) 2020-2026 Paweł Kuna）；码点表 codepoints.json；中文语义表 semantic.json。

用法：
  python inject_icons.py deck.pptx icons.json --colors "INK=111111,ACCENT=E4002B,MUTED=6E6E6E"
  python inject_icons.py deck.pptx icons.json --replace     # 先清掉已有图标再注入
  python inject_icons.py deck.pptx icons.json --no-verify

规格格式（中文键；图标=中文语义名或 Tabler 原名；锚点=施工时的 objectName）：
  {
    "pages": {
      "2": [
        {"图标": "安全", "锚点": "能力卡一", "位置": "上方居中", "颜色": "ACCENT",
         "尺寸": 0.55, "间距": 0.15},
        {"图标": "rocket", "锚点": "能力卡二", "位置": "内部左上", "颜色": "INK", "尺寸": 0.4}
      ]
    }
  }
  位置：上方居中 / 下方居中 / 左侧居中 / 右侧居中 / 内部左上 / 内部右上 / 内部居中
  颜色：--colors 里的角色名，或直接六位十六进制
  尺寸：图标视觉外框边长（英寸），按最长边归一化

纪律（规格书里要声明图标政策）：家族固定、一律单色、每页 ≤3 个、禁 emoji、
禁「图标 + 圆角卡片 + 投影」三件套。qa.py 会对每页图标数与有效 PPI 兜底。

验证：重开计数（icon:* 命名）+ COM 兼容打开 + 逐条有效 PPI。
防重复：插图像统一命名 icon:中文名；重复注入会被拦，需 --replace。
退出码：0 成功；1 验证不一致；2 跳过验证；规格错误直接报错、文件不动。
"""
import json
import os
import shutil
import sys
import tempfile

ICON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "icons")
POSITIONS = ("上方居中", "下方居中", "左侧居中", "右侧居中",
             "内部上方居中", "内部下方居中",
             "内部左上", "内部右上", "内部左下", "内部右下", "内部居中")

def _assets():
    d = os.path.abspath(ICON_DIR)
    font = os.path.join(d, "tabler-icons.ttf")
    if not os.path.exists(font):
        sys.exit(f"找不到图标字体 {font}——技能安装不完整（assets/icons 缺失）")
    sem = json.load(open(os.path.join(d, "semantic.json"), encoding="utf-8"))
    cps = json.load(open(os.path.join(d, "codepoints.json"), encoding="utf-8"))
    return font, sem, cps

def render_icon(font_path, cp_hex, color, size_in, dpi, out):
    """渲染单个图标：字形紧裁 → 按最长边归一到 size_in×dpi → 透明 PNG。"""
    from PIL import Image, ImageDraw, ImageFont
    px = max(int(round(size_in * dpi)), 8)
    big = px * 2
    font = ImageFont.truetype(font_path, px)
    img = Image.new("RGBA", (big, big), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    ch = chr(int(cp_hex, 16))
    d.text((big / 2, big / 2), ch, font=font, fill="#" + color, anchor="mm")
    bbox = img.getbbox()
    if not bbox:
        raise RuntimeError("字形渲染为空——字体里没有这个码点")
    glyph = img.crop(bbox)
    w, h = glyph.size
    scale = px / max(w, h)
    glyph = glyph.resize((max(1, int(round(w * scale))), max(1, int(round(h * scale)))),
                         Image.LANCZOS)
    canvas = Image.new("RGBA", (px, px), (0, 0, 0, 0))
    canvas.paste(glyph, ((px - glyph.width) // 2, (px - glyph.height) // 2), glyph)
    canvas.save(out)
    return px

def _place(sh, icon_w, icon_h, pos, gap_emu):
    """锚点形状 + 位置 -> (left, top) EMU。全程 EMU 运算（混单位会让图跑出画布）。"""
    cx = sh.left + (sh.width - icon_w) // 2
    cy = sh.top + (sh.height - icon_h) // 2
    if pos == "上方居中":
        return cx, sh.top - gap_emu - icon_h
    if pos == "下方居中":
        return cx, sh.top + sh.height + gap_emu
    if pos == "左侧居中":
        return sh.left - gap_emu - icon_w, cy
    if pos == "右侧居中":
        return sh.left + sh.width + gap_emu, cy
    if pos == "内部上方居中":
        return cx, sh.top + gap_emu                     # 卡片上区 + 标题在下：最常用的成组摆法
    if pos == "内部下方居中":
        return cx, sh.top + sh.height - gap_emu - icon_h
    if pos == "内部左上":
        return sh.left + gap_emu, sh.top + gap_emu
    if pos == "内部右上":
        return sh.left + sh.width - gap_emu - icon_w, sh.top + gap_emu
    if pos == "内部左下":
        return sh.left + gap_emu, sh.top + sh.height - gap_emu - icon_h
    if pos == "内部右下":
        return sh.left + sh.width - gap_emu - icon_w, sh.top + sh.height - gap_emu - icon_h
    return cx, cy                                     # 内部居中

def parse_colors(arg):
    out = {}
    for part in (arg or "").split(","):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip().lstrip("#").upper()
    return out

def inject(src, spec, colors, dpi, replace):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("缺依赖：python -m pip install python-pptx")
    font_path, sem, cps = _assets()
    jobs, errors = [], []
    prs = Presentation(src)
    for si, slide in enumerate(prs.slides, 1):
        items = spec.get("pages", {}).get(str(si), [])
        if not items:
            continue
        by_name = {}
        for sh in slide.shapes:
            if sh.name:
                by_name[sh.name] = sh
        for it in items:
            icon = it.get("图标", "")
            tname = sem.get(icon, icon)
            if tname not in cps:
                errors.append(f"S{si}: 图标「{icon}」不在语义表里；可用："
                              + "、".join(sorted(sem))[:120] + "…或直接给 Tabler 原名")
                continue
            anchor = it.get("锚点")
            if anchor not in by_name:
                errors.append(f"S{si}: 锚点形状「{anchor}」不存在；本页有："
                              + "、".join(sorted(n for n in by_name if n))[:120])
                continue
            pos = it.get("位置", "上方居中")
            if pos not in POSITIONS:
                errors.append(f"S{si}: 位置「{pos}」不合法；可选：{'、'.join(POSITIONS)}")
                continue
            raw = str(it.get("颜色", "INK")).lstrip("#").upper()
            color = colors.get(raw, raw if len(raw) == 6 else None)
            if not color:
                errors.append(f"S{si}: 颜色「{it.get('颜色')}」既不是 --colors 里的角色"
                              f"（{ '、'.join(colors) or '空' }）也不是六位十六进制")
                continue
            size = float(it.get("尺寸", 0.5))
            jobs.append((si, slide, by_name[anchor], icon, tname, pos, color, size,
                         float(it.get("间距", 0.12))))
    if errors:
        for e in errors:
            print("✗ " + e)
        sys.exit("图标规格有误（文件保持原样）——修正后重跑")

    # 防重复：已存在 icon:* 图片时先拦（--replace 则先清）
    stale = [(si, sh) for si, slide in enumerate(prs.slides, 1) for sh in slide.shapes
             if sh.name and sh.name.startswith("icon:")]
    if stale and not replace:
        sys.exit(f"文件里已有 {len(stale)} 个图标——改规格后重注入请加 --replace")
    for _si, sh in stale:
        sh._element.getparent().remove(sh._element)

    if not jobs:
        print("规格里没有可注入的图标，文件未改动")
        return 0, 0

    tmpdir = tempfile.mkdtemp(prefix="icons_")
    placed = 0
    try:
        for i, (si, slide, anchor, icon, tname, pos, color, size, gap) in enumerate(jobs):
            png = os.path.join(tmpdir, f"i{i}.png")
            try:
                px = render_icon(font_path, cps[tname], color, size, dpi, png)
            except RuntimeError as e:
                errors.append(f"S{si}: 图标「{icon}」({tname}) 渲染失败——{e}")
                continue
            box = Emu(int(size * 914400))
            gap_emu = int(gap * 914400)
            left, top = _place(anchor, box, box, pos, gap_emu)
            left, top = max(int(left), 0), max(int(top), 0)
            pic = slide.shapes.add_picture(png, left, top, width=box, height=box)
            pic._element.nvPicPr.cNvPr.set("name", f"icon:{icon}")
            placed += 1
            print(f"  S{si}：{icon}({tname}) {pos} @ 锚点「{anchor.name}」"
                  f" {size:g}\" #{color} → {px}px")
        if errors:  # 不写半个文件：有失败就整份保持原样
            for e in errors:
                print("✗ " + e)
            sys.exit("图标规格有误（文件保持原样）——修正后重跑")
        prs.save(src)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    print(f"已插入 {placed} 个图标（{dpi}dpi / Tabler 线性）")
    return placed, len(jobs)

def verify(src, placed, colors):
    from pptx import Presentation
    prs = Presentation(src)
    icons = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.name and sh.name.startswith("icon:"):
                icons.append((si, sh))
    if len(icons) != placed:
        print(f"  ✗ 重开计数 {len(icons)} ≠ 注入 {placed}")
        return False
    for si, sh in icons:
        ppi = round(sh.image.size[0] / (sh.width / 914400))
        print(f"  ✓ S{si} {sh.name} 有效 PPI ≈ {ppi}")
    try:
        import win32com.client
        app = win32com.client.dynamic.Dispatch("PowerPoint.Application")
        try:
            app.DisplayAlerts = 1
        except Exception:
            pass
        pres = app.Presentations.Open(os.path.abspath(src), ReadOnly=True, WithWindow=False)
        pres.Close()
        try:
            if app.Presentations.Count == 0:
                app.Quit()
        except Exception:
            pass
        print(f"  ✓ PowerPoint 兼容 COM 打开无修复"
              f"（{app.Name} {app.Version}.{app.Build}）")
    except ImportError:
        print("  ⚠ 无 COM，跳过打开体检")
    except Exception as e:
        print(f"  ✗ COM 打开失败：{str(e)[:110]}")
        return False
    return True

def main():
    # 规范参数扫描：带值旗标要吃掉它的值，否则 "--colors INK=111" 的值会被当成位置参数
    # 未知旗标/给错的值一律报错退出——静默退回默认值会让 AI 以为设置已生效。
    argv = sys.argv[1:]
    known = {"--colors", "--dpi", "--no-verify", "--replace"}
    for a in argv:
        if not a.startswith("--"):
            continue
        base, eq, _ = a.partition("=")
        if base not in known:
            sys.exit(f"未知参数 {base}；可用：{' '.join(sorted(known))}")
        if eq and base in ("--colors", "--dpi"):
            sys.exit(f"{base} 不支持 = 形式（--colors 的值本身含 =），请用空格：{base} <值>")
    pos, colors_arg, dpi, flags, i = [], "", 800, set(), 0
    while i < len(argv):
        a = argv[i]
        if a in ("--colors", "--dpi"):
            if i + 1 >= len(argv):
                sys.exit(f"{a} 缺少值")
            val = argv[i + 1]
            if a == "--colors":
                colors_arg = val
            else:
                try:
                    dpi = int(val)
                except ValueError:
                    sys.exit(f"{a} 的值 {val!r} 不是整数")
            i += 2
            continue
        if a.startswith("--"):
            flags.add(a)
            i += 1
            continue
        pos.append(a)
        i += 1
    if len(pos) != 2:
        sys.exit(__doc__)
    src, spec_path = pos
    if not os.path.exists(src):
        sys.exit("找不到文件：" + src)
    try:
        spec = json.load(open(spec_path, encoding="utf-8"))
    except Exception as e:
        sys.exit(f"规格 {spec_path} 读不了：{e}")
    colors = parse_colors(colors_arg)
    placed, _ = inject(src, spec, colors, dpi, replace="--replace" in flags)
    if not placed:
        return 0
    if "--no-verify" in flags:
        print("⚠ 已插入但跳过验证（--no-verify）")
        return 2
    return 0 if verify(src, placed, colors) else 1

if __name__ == "__main__":
    sys.exit(main())
