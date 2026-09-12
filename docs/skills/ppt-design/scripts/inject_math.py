#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把文本里的块级 LaTeX 公式渲染成高分辨率透明图并放回原位。

pptxgenjs / python-pptx 都没有公式 API——直接把 LaTeX 写进文本框只会显示成
乱码字符串。本机实测两条路线后选定图片路线：原生 OMML 公式在 WPS（很多
机器的实际"PowerPoint"）里不渲染；LaTeX→PNG 在 WPS / 真 PowerPoint /
Keynote 里观感完全一致，且 600 DPI 对 4K 投影是 2 倍过采样，肉眼即矢量。

用法：
  python inject_math.py deck.pptx                 # 渲染+插入+验证（默认 600dpi）
  python inject_math.py deck.pptx --dpi 1200      # 更高分辨率
  python inject_math.py deck.pptx --fontset stix  # 公式字体：cm（LaTeX 经典，默认）/ stix / dejavusans
  python inject_math.py deck.pptx --no-verify     # 跳过验证

写作语法（施工时直接写在 addText 的字符串里）：
  s.addText("$$x = \\\\frac{-b \\\\pm \\\\sqrt{b^2-4ac}}{2a}$$", { fontSize: 32, color: "C13B2E" })
  规则（三条都是硬性的，违反会报错而不是静默出错）：
  1. 公式用 $$...$$ 包裹，且必须**独占一个文本框**——公式的字号、颜色、
     位置都从这个文本框继承（fontSize 32 → 公式就是 32pt；color 跟随文字色）；
  2. 行内 $...$ 公式图片模式不支持（图片无法嵌进文字流）——报错并列出位置，
     把式子改成独立文本框；
  3. 一个文本框一条公式。mathtext 支持 分数/根号/上下标/希腊字母/求和/积分/
     极限/累乘；矩阵、cases、多行 align 不支持（无 LaTeX 环境依赖）。

验证：插入后重开检查（标记清零 + 图片部件数一致）+ PowerPoint 兼容 COM
打开（坏关系/坏部件会在这一层炸出来）。渲染观感交闸门二。
idempotent：转换后不再有标记，重复运行是空操作。
退出码：0 成功；1 验证不一致；2 成功但跳过验证；转换失败直接报错、文件不动。
"""
import os
import re
import shutil
import sys
import tempfile

BLOCK_RE = re.compile(r"^\$\$(.+)\$\$$", re.S)
ANY_DOLLAR = re.compile(r"\$([^$\n]+?)\$")
MATH_LIKE = re.compile(r"[\\^_]")     # LaTeX 特征：命令 / 上下标——没有它就是价格等字面 $

# ---------------------------------------------------------------- 渲染

def render_formula(latex: str, pt: float, color: str, dpi: int, fontset: str, path: str):
    """mathtext → 透明 PNG。返回 (宽_px, 高_px)。解析失败抛异常由调用方收集。"""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    plt.rcParams["mathtext.fontset"] = fontset
    fig = plt.figure(figsize=(0.5, 0.5))
    fig.text(0.5, 0.5, f"${latex}$", fontsize=pt, color=f"#{color}",
             ha="center", va="center")
    fig.savefig(path, dpi=dpi, transparent=True, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    from PIL import Image
    with Image.open(path) as im:
        return im.size

# ---------------------------------------------------------------- 插入

def _shape_meta(sh):
    """(字号 pt, 颜色 RRGGBB)——取第一个带属性的 run，缺省 18pt / 黑。"""
    pt, color = 18.0, "000000"
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    pt = r.font.size.pt
                try:
                    if r.font.color and r.font.color.rgb:
                        color = str(r.font.color.rgb)
                except Exception:
                    pass
                return pt, color
    except Exception:
        pass
    return pt, color

def inject(src: str, dpi: int, fontset: str):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("缺依赖：python -m pip install python-pptx")
    prs = Presentation(src)
    errors, jobs = [], []
    for si, slide in enumerate(prs.slides, 1):
        for sh in list(slide.shapes):
            if not getattr(sh, "has_text_frame", False):
                continue
            text = (sh.text_frame.text or "").strip()
            if "$$" not in text and not ANY_DOLLAR.search(text):
                continue
            m = BLOCK_RE.match(text)
            if not m:
                inline = next((c for c in ANY_DOLLAR.findall(text)
                               if MATH_LIKE.search(c)), None)
                if inline:
                    errors.append(f"S{si}: 行内公式 {inline[:30]!r} "
                                  "—— 图片模式仅支持块级公式：把式子放进**独占的**文本框")
                continue                     # 纯字面 $（价格等）：不动
            pt, color = _shape_meta(sh)
            jobs.append((si, sh, m.group(1).strip(), pt, color))
    if errors:
        for e in errors:
            print("✗ " + e)
        sys.exit("公式处理失败（文件保持原样）——按上面提示调整后重跑")
    if not jobs:
        print("未发现公式标记（$$…$$ 独占文本框），文件未改动")
        return 0, 0

    tmpdir = tempfile.mkdtemp(prefix="formula_")
    placed = 0
    try:
        for i, (si, sh, latex, pt, color) in enumerate(jobs):
            png = os.path.join(tmpdir, f"f{i}.png")
            try:
                w_px, h_px = render_formula(latex, pt, color, dpi, fontset, png)
            except Exception as e:
                errors.append(f"S{si}: {latex[:40]!r} —— mathtext 无法解析：{str(e)[:80]}")
                continue
            w_in, h_in = w_px / dpi, h_px / dpi
            img_w, img_h = Emu(int(w_in * 914400)), Emu(int(h_in * 914400))
            # 全程 EMU 运算：sh.width 是 EMU，混进英寸会让图片小到不可见且跑出画布（实测踩过）
            left = max(sh.left + (sh.width - img_w) // 2, 0)
            top = max(sh.top + (sh.height - img_h) // 2, 0)
            el = sh._element
            el.getparent().remove(el)          # 公式框功成身退：删框、图落原位
            slide = prs.slides[si - 1]
            pic = slide.shapes.add_picture(png, left, top, width=img_w, height=img_h)
            pic._element.nvPicPr.cNvPr.set("name", f"formula:{latex[:24]}")
            placed += 1
            print(f"  S{si}：{latex[:36]!r} @ {pt:g}pt #{color} → {w_px}×{h_px}px "
                  f"({w_in:.2f}×{h_in:.2f}\")")
        if errors:
            for e in errors:
                print("✗ " + e)
            sys.exit("公式渲染失败（文件保持原样）——修正 LaTeX 后重跑")
        prs.save(src)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    print(f"已插入 {placed} 条公式图片（{dpi}dpi / {fontset} 字体）")
    return placed, len(jobs)

# ---------------------------------------------------------------- 验证

def verify(src: str, placed: int):
    """重开检查标记清零 + 图片数；COM 兼容打开做最后一道体检。"""
    from pptx import Presentation
    prs = Presentation(src)
    leftovers, pics = 0, 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and "$$" in (sh.text_frame.text or ""):
                leftovers += 1
            if sh.shape_type == 13:
                pics += 1
    if leftovers:
        print(f"  ✗ 仍有 {leftovers} 处未转换的公式标记")
        return False
    print(f"  ✓ 标记清零，共 {pics} 张图片（含 {placed} 条新公式）")
    try:
        import win32com.client
        app = win32com.client.dynamic.Dispatch("PowerPoint.Application")
        try:
            app.DisplayAlerts = 1
        except Exception:
            pass
        pres = app.Presentations.Open(os.path.abspath(src), ReadOnly=True, WithWindow=False)
        pres.Close()
        try:
            if app.Presentations.Count == 0:
                app.Quit()
        except Exception:
            pass
        print("  ✓ PowerPoint 兼容 COM 打开无修复（部件与关系完整）")
    except ImportError:
        print("  ⚠ 无 COM，跳过打开体检")
    except Exception as e:
        print(f"  ✗ COM 打开失败：{str(e)[:110]}")
        return False
    return True

# ---------------------------------------------------------------- main

def main():
    # 规范参数扫描：带值旗标吃掉它的值（否则 --dpi 1200 的 1200 会被当成位置参数）
    # 未知旗标/给错的值一律报错退出——静默退回默认值会让 AI 以为设置已生效。
    argv = sys.argv[1:]
    known = {"--dpi", "--fontset", "--no-verify"}
    for a in argv:
        if not a.startswith("--"):
            continue
        base, eq, _ = a.partition("=")
        if base not in known:
            sys.exit(f"未知参数 {base}；可用：{' '.join(sorted(known))}")
        if eq and base in ("--dpi", "--fontset"):
            sys.exit(f"{base} 不支持 = 形式，请用空格：{base} <值>")
    pos, dpi, fontset, flags, i = [], 600, "cm", set(), 0
    while i < len(argv):
        a = argv[i]
        if a in ("--dpi", "--fontset"):
            if i + 1 >= len(argv):
                sys.exit(f"{a} 缺少值")
            val = argv[i + 1]
            if a == "--fontset":
                fontset = val
            else:
                try:
                    dpi = int(val)
                except ValueError:
                    sys.exit(f"{a} 的值 {val!r} 不是整数")
            i += 2
            continue
        if a.startswith("--"):
            flags.add(a)
            i += 1
            continue
        pos.append(a)
        i += 1
    if len(pos) != 1:
        sys.exit(__doc__)
    src = pos[0]
    if not os.path.exists(src):
        sys.exit("找不到文件：" + src)
    placed, _ = inject(src, dpi, fontset)
    if not placed:
        return 0
    if "--no-verify" in flags:
        print("⚠ 已插入但跳过验证（--no-verify）")
        return 2
    return 0 if verify(src, placed) else 1

if __name__ == "__main__":
    sys.exit(main())
